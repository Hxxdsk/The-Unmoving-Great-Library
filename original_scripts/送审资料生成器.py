# -*- coding: utf-8 -*-
"""
送审资料生成器
"""
import re
import os
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side  # 导入填充/字体/对齐/边框样式
# ============================================================
# 1. 固定项 / 规则常量
# ============================================================
OUTPUT_HEADERS = [
    "开发款号(*)","商品类型(*)","是否底胚(*)","商品名称(*)","年份(*)","季节(*)",
    "大类名称(*)","品类编码(*)","小类编码(*)","性别(*)","尺码组(*)","颜色编码(*)",
    "上下装(*)","材质","面料名称(*)","面料成分(*)","里料成分(*)","填充物成分(*)",
    "安全技术级别(*)","执行标准(*)","包材单价(*)","标费档次(*)","成本价(*)","标准价(*)",
    "吊牌价(*)","生产厂商(*)","渠道(*)","平台(*)","商品备注","辅料(*)",
]
CLASS_RULES = [
    ("外套/夹克",     ["外套", "夹克"],        "FZ/T 81007-2022", "GB 18401-2010 C类"),
    ("衬衫",          ["衬衫"],                "FZ/T 81007-2022", "GB 18401-2010 B类"),
    ("休闲裤",        ["裤"],              "FZ/T 81007-2022", "GB 18401-2010 B类"),
    ("毛衣（针织衫）", ["毛衣", "针织衫"],      "FZ/T 73034-2021", "GB 18401-2010 B类"),
    ("T恤",           ["t恤"],                 "GB/T 22849-2024", "GB 18401-2010 B类"),
    ("卫衣",          ["卫衣"],                "GB/T 22849-2024", "GB 18401-2010 B类"),
]
# 新增：黄色填充样式
YELLOW_FILL = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
# 新增：表头样式（严格按照示例模板格式：Calibri 11 加粗、灰色底纹、居中、细边框）
HEADER_FILL = PatternFill(start_color="FFE0E0E0", end_color="FFE0E0E0", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True)
HEADER_ALIGN = Alignment(horizontal="center")
HEADER_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)

def normalize_key(s):
    if s is None:
        return ""
    return re.sub(r"\s+", "", str(s)).lower()
# ----------------------【面料成分识别 / 标准化】----------------------
FABRIC_LABELS = {
    "面料成分", "面料", "里料成分", "里料", "底料成分", "底料",
    "填充物成分", "填充物", "内里", "配料", "辅料",
}
FABRIC_LABEL_PATTERN = r'(?:面料成分|里料成分|底料成分|填充物成分|面料|里料|底料|填充物|内里|配料|辅料)\s*[:：]?'


def _format_percent(value):
    """把 98 / 98.0 / 98% / 98％ 统一为 98%，小数按有效精度保留。"""
    s = str(value).strip().replace("％", "%").replace("%", "").strip()
    try:
        num = float(s)
    except (TypeError, ValueError):
        return ""
    if num.is_integer():
        return "%d%%" % int(num)
    return ("%g" % num) + "%"


def _clean_material_name(name):
    """只清理材质名称两侧标点；不再删除“面料/里料/底料”等标签。"""
    s = str(name or "").strip()
    s = re.sub(r'^[·•\-.,，。;；]+|[·•\-.,，。;；]+$', '', s).strip()
    return s


def _normalize_fabric_label(token):
    """标签原义保留，仅压缩冒号两侧空格。"""
    s = str(token or "").strip()
    m = re.fullmatch(r'(.+?)\s*([:：])?', s)
    if not m:
        return s
    name = m.group(1).strip()
    colon = m.group(2) or ""
    return name + colon


def _is_fabric_label(token):
    s = str(token or "").strip()
    s = re.sub(r'\s*[:：]\s*$', '', s).strip()
    return s in FABRIC_LABELS


def _tokenize_fabric(text):
    """把面料文本拆成标签(L)、百分比(N)、普通文字(T)，保留位置用于合理配对。"""
    if text is None:
        return []
    s = str(text).strip().replace("％", "%")
    if not s:
        return []

    # 常见成分分隔符统一为空格；冒号保留，用于保存“面料:/里料:/底料:”等标签。
    s = re.sub(r'[;；,，、/|+\n\r\t]+', ' ', s)
    s = re.sub(r'\s+', ' ', s).strip()

    pattern = re.compile(
        FABRIC_LABEL_PATTERN
        + r'|\d+(?:\.\d+)?\s*%'
        + r'|(?<!\d)\d+(?:\.\d+)?(?!\d)'
        + r'|[^\d%\s]+'
    )

    tokens = []
    for m in pattern.finditer(s):
        raw = m.group(0).strip()
        if not raw:
            continue
        if _is_fabric_label(raw):
            kind = "L"
            value = _normalize_fabric_label(raw)
        elif re.fullmatch(r'\d+(?:\.\d+)?\s*%?', raw):
            kind = "N"
            value = raw
        else:
            kind = "T"
            value = _clean_material_name(raw)
            if not value:
                continue
        tokens.append({
            "kind": kind,
            "value": value,
            "start": m.start(),
            "end": m.end(),
        })
    return tokens


def _pair_fabric_tokens(tokens):
    """
    将百分比与左右相邻材质配对。
    - 标签(L)绝不参与配对，避免“里料:100%聚酯纤维”误成“100%里料”。
    - 百分比左右都有文字时，优先选择距离更近的一侧；距离相同优先左侧，兼容“聚酯纤维98%”。
    """
    used = set()
    pairs = []

    for i, tok in enumerate(tokens):
        if tok["kind"] != "N" or i in used:
            continue

        candidates = []
        if i - 1 >= 0 and (i - 1) not in used and tokens[i - 1]["kind"] == "T":
            left = tokens[i - 1]
            if not left["value"].endswith((":", "：")):
                gap = max(0, tok["start"] - left["end"])
                candidates.append((gap, 0, i - 1))  # 同距离时优先左侧

        if i + 1 < len(tokens) and (i + 1) not in used and tokens[i + 1]["kind"] == "T":
            right = tokens[i + 1]
            if not right["value"].endswith((":", "：")):
                gap = max(0, right["start"] - tok["end"])
                candidates.append((gap, 1, i + 1))

        if not candidates:
            continue

        _, _, text_idx = min(candidates, key=lambda x: (x[0], x[1]))
        pct = _format_percent(tok["value"])
        name = _clean_material_name(tokens[text_idx]["value"])
        if not pct or not name or _is_fabric_label(name):
            continue

        used.add(i)
        used.add(text_idx)
        pairs.append({
            "pct": pct,
            "name": name,
            "start_idx": min(i, text_idx),
            "end_idx": max(i, text_idx),
            "num_idx": i,
            "text_idx": text_idx,
        })

    pairs.sort(key=lambda x: x["start_idx"])
    return pairs, used


def parse_fabric_components(text):
    """
    识别真实“百分比 + 材质”成分对，标签不会被当成材质。

    示例：
      98%聚酯纤维 2%棉                 -> 98%聚酯纤维 / 2%棉
      聚酯纤维98% 棉2%                 -> 98%聚酯纤维 / 2%棉
      面料:98%聚酯纤维 里料:100%棉     -> 98%聚酯纤维 / 100%棉
      里料:100%聚酯纤维                -> 100%聚酯纤维（不会识别成100%里料）
    """
    tokens = _tokenize_fabric(text)
    pairs, _ = _pair_fabric_tokens(tokens)
    return [(p["pct"], p["name"]) for p in pairs]


def _join_fabric_parts(parts):
    """标签后的成分紧跟冒号，其余项目用单空格分隔。"""
    result = ""
    for part in (p for p in parts if p):
        if not result:
            result = part
        elif result.endswith((":", "：")):
            result += part
        else:
            result += " " + part
    return result.strip()


def clean_fabric_text(text):
    """
    标准化面料文本：
    1) 成分统一为“百分比%材质名”；
    2) 面料/里料/底料/填充物等标签原样保留，不参与成分配对；
    3) 无法识别为成分的其它文字也保留，不静默删除。
    """
    if text is None:
        return ""
    raw = str(text).strip()
    if not raw:
        return ""

    tokens = _tokenize_fabric(raw)
    if not tokens:
        return raw

    pairs, used = _pair_fabric_tokens(tokens)
    pair_by_start = {p["start_idx"]: p for p in pairs}
    pair_member = set()
    for p in pairs:
        pair_member.add(p["num_idx"])
        pair_member.add(p["text_idx"])

    parts = []
    i = 0
    while i < len(tokens):
        if i in pair_by_start:
            p = pair_by_start[i]
            parts.append(p["pct"] + p["name"])
            i = p["end_idx"] + 1
            continue

        if i in pair_member:
            i += 1
            continue

        tok = tokens[i]
        if tok["kind"] == "L":
            parts.append(_normalize_fabric_label(tok["value"]))
        elif tok["kind"] == "N":
            # 单独出现的百分比也不删除。
            parts.append(_format_percent(tok["value"]) or tok["value"])
        else:
            parts.append(tok["value"])
        i += 1

    return _join_fabric_parts(parts)

def detect_class(style):
    s = normalize_key(style)
    candidates = []
    for name, kws, std, sec in CLASS_RULES:
        for kw in kws:
            if kw in s:
                candidates.append((len(kw), name, std, sec))
    if not candidates:
        return None
    candidates.sort(key=lambda x: -x[0])
    return candidates[0][1], candidates[0][2], candidates[0][3]
# ============================================================
# 2. 主数据加载
# ============================================================
def load_master(master_path):
    if not master_path or not os.path.exists(master_path):
        raise FileNotFoundError("找不到主数据表：%s" % master_path)
    wb = openpyxl.load_workbook(master_path, data_only=True)
    if "新品类层级编码及标费档次" not in wb.sheetnames:
        raise KeyError("主数据表缺少【新品类层级编码及标费档次】工作表")
    ws = wb["新品类层级编码及标费档次"]
    cat_map, price_map = {}, {}
    for r in range(3, ws.max_row + 1):
        c = ws.cell(r, 3).value
        b = ws.cell(r, 2).value
        d = ws.cell(r, 4).value
        if c and b:
            k = normalize_key(c)
            if k not in cat_map:
                cat_map[k] = (str(b).strip(), str(d).strip() if d is not None else "")
        h = ws.cell(r, 8).value
        j = ws.cell(r, 10).value
        l = ws.cell(r, 12).value
        if h and j is not None:
            k = normalize_key(h)
            if k not in price_map:
                price_map[k] = (j, l)
    if "尺码表" not in wb.sheetnames:
        raise KeyError("主数据表缺少【尺码表】工作表")
    ws2 = wb["尺码表"]
    size_map = {}
    for r in range(2, ws2.max_row + 1):
        a = ws2.cell(r, 1).value
        b = ws2.cell(r, 2).value
        if a is not None and b:
            size_map[normalize_key(b)] = str(a).strip()
    if "颜色" not in wb.sheetnames:
        raise KeyError("主数据表缺少【颜色】工作表")
    ws3 = wb["颜色"]
    color_map = {}
    for r in range(2, ws3.max_row + 1):
        a = ws3.cell(r, 1).value
        b = ws3.cell(r, 2).value
        if a and b:
            color_map[str(a).strip()] = str(b).strip()
    return cat_map, price_map, size_map, color_map
def get_full_color_list(master_path):
    """读取全部颜色库，返回list[(颜色名,编码)]供搜索弹窗使用"""
    if not master_path or not os.path.exists(master_path):
        raise FileNotFoundError("找不到主数据表：%s" % master_path)
    wb = openpyxl.load_workbook(master_path, data_only=True)
    if "颜色" not in wb.sheetnames:
        raise KeyError("主数据表缺少【颜色】工作表")
    ws3 = wb["颜色"]
    full_list = []
    for r in range(2, ws3.max_row +1):
        name = ws3.cell(r,1).value
        code = ws3.cell(r,2).value
        if name and code:
            full_list.append((str(name).strip(), str(code).strip()))
    return full_list
# ============================================================
# 3. 字段处理函数
# ============================================================
def parse_cost(value):
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).replace(",", "").strip())
    except ValueError:
        return None
def extract_first_material(fabric):
    """提取第一组成分的真实材质名；“面料/里料/底料”等标签不作为材质。"""
    if fabric is None or fabric == "":
        return ""
    components = parse_fabric_components(fabric)
    if components:
        return components[0][1]

    tokens = _tokenize_fabric(fabric)
    for tok in tokens:
        if tok["kind"] == "T" and not _is_fabric_label(tok["value"]):
            return tok["value"]
    return ""
def suggest_color(color, color_map):
    def score(a, b):
        if not a or not b:
            return 0.0
        sa, sb = set(a), set(b)
        base = 2.0 * len(sa & sb) / (len(sa) + len(sb))
        if a in b or b in a:
            base += 1.0
        return base
    ranked = sorted(color_map.keys(), key=lambda n: score(color, n), reverse=True)
    return [(n, color_map[n]) for n in ranked[:3] if score(color, n) > 0.1]
def split_colors(text):
    if not text:
        return []
    parts = re.split(r"[\s,;，；、/]+", str(text).strip())
    seen, out = set(), []
    for p in parts:
        p = p.strip()
        if p and p not in seen:
            seen.add(p)
            out.append(p)
    return out
# ============================================================
# 4. 单行处理
# ============================================================
def process_row(row, cfg, warnings):
    cat_map, price_map, size_map, color_map = cfg["maps"]
    st = cfg["settings"]
    missing_colors = cfg.get("missing_colors")
    color_issues = cfg.get("color_issues")
    style_raw = str(row.get("款式", "") or "").strip()
    style_key = normalize_key(style_raw)
    no_style = not style_key
    款号 = str(row.get("款号", "") or "").strip()
    out = {}
    out["开发款号(*)"] = st["开发前缀"] + 款号
    out["商品类型(*)"] = "公版"
    out["是否底胚(*)"] = "否"
    out["商品名称(*)"] = st["产品前缀"] + style_raw
    out["年份(*)"] = st["年份"]
    out["季节(*)"] = st["季节"]
    out["大类名称(*)"] = st["大类"]
    out["性别(*)"] = st["性别"]
    if no_style:
        out["品类编码(*)"] = ""
        out["小类编码(*)"] = ""
        if 款号:
            warnings.append("款号「%s」的款式为空，无法匹配品类/小类/标准价/标费档次。" % 款号)
    elif style_key in cat_map:
        b, d = cat_map[style_key]
        out["品类编码(*)"] = b
        out["小类编码(*)"] = d
    else:
        out["品类编码(*)"] = ""
        out["小类编码(*)"] = ""
        warnings.append("款式「%s」未匹配到【新品类层级编码及标费档次】C列（小类名称），品类编码/小类编码留空。" % style_raw)
    尺码 = str(row.get("尺码", "") or "").strip()
    if not 尺码:
        out["尺码组(*)"] = ""
    elif normalize_key(尺码) in size_map:
        out["尺码组(*)"] = size_map[normalize_key(尺码)]
    else:
        out["尺码组(*)"] = ""
        warnings.append("尺码「%s」未匹配到【尺码表】B列（尺码组）。" % 尺码)
    colors = split_colors(row.get("颜色", ""))
    codes, missing = [], []
    matched_pairs = []
    for col in colors:
        if col in color_map:
            code = color_map[col]
            codes.append(code)
            matched_pairs.append((col, code))
        else:
            missing.append(col)
    out["颜色编码(*)"] = ";".join(codes)

    # 颜色完整性检测：原始颜色数量、已匹配代码数量、唯一代码数量三者应一致。
    # 若两个不同颜色指向同一个编码，也明确提示具体款号/款式。
    duplicate_codes = {}
    code_to_colors = {}
    for col, code in matched_pairs:
        code_to_colors.setdefault(str(code), []).append(col)
    for code, cols in code_to_colors.items():
        if len(cols) > 1:
            duplicate_codes[code] = cols

    expected_count = len(colors)
    matched_count = len(codes)
    unique_code_count = len(set(str(c) for c in codes))
    if expected_count and (matched_count != expected_count or unique_code_count != expected_count):
        issue = {
            "款号": 款号,
            "款式": style_raw,
            "原始颜色": str(row.get("颜色", "") or "").strip(),
            "颜色列表": colors,
            "原颜色数": expected_count,
            "匹配代码数": matched_count,
            "唯一代码数": unique_code_count,
            "缺失颜色": missing[:],
            "重复代码": duplicate_codes,
        }
        if color_issues is not None:
            color_issues.append(issue)

        reason_parts = []
        if missing:
            reason_parts.append("未匹配：%s" % "、".join(missing))
        if duplicate_codes:
            dup_text = "；".join(
                "%s → %s" % (code, "、".join(cols))
                for code, cols in duplicate_codes.items()
            )
            reason_parts.append("不同颜色共用编码：%s" % dup_text)
        if not reason_parts:
            reason_parts.append("颜色数与编码数不一致")
        warnings.append(
            "颜色检测异常｜款号「%s」款式「%s」：原有%d个颜色，匹配%d个代码，唯一代码%d个；%s。"
            % (款号 or "未填写", style_raw or "未填写", expected_count, matched_count,
               unique_code_count, "；".join(reason_parts))
        )

    if missing:
        for mcol in missing:
            top = suggest_color(mcol, color_map)
            if missing_colors is not None:
                missing_colors.append({"orig": mcol, "sugs": top})
            if top:
                warnings.append(
                    "颜色「%s」未匹配，推荐：%s。"
                    % (mcol, " / ".join("%s(%s)" % (n, c) for n, c in top))
                )
            else:
                warnings.append("颜色「%s」未匹配，且无相近颜色推荐。" % mcol)
    out["上下装(*)"] = str(row.get("上下装", "") or "").strip()
    out["材质"] = ""
    # ----------------【面料处理修改点】先清洗原始面料文本 ----------------
    fabric_raw = str(row.get("面料", "") or "")
    fabric_clean = clean_fabric_text(fabric_raw)
    out["面料名称(*)"] = extract_first_material(fabric_clean)
    out["面料成分(*)"] = fabric_clean
    out["里料成分(*)"] = "无"
    out["填充物成分(*)"] = "无"
    cls = detect_class(style_raw)
    if cls is None:
        out["执行标准(*)"] = ""
        out["安全技术级别(*)"] = ""
        if style_raw:
            warnings.append("款式「%s」未能识别到类目，执行标准/安全技术级别留空。" % style_raw)
    else:
        _, std, sec = cls
        out["执行标准(*)"] = std
        out["安全技术级别(*)"] = sec
    out["包材单价(*)"] = st["包材单价"]
    if no_style:
        out["标费档次(*)"] = ""
        out["标准价(*)"] = ""
    elif style_key in price_map:
        j, l = price_map[style_key]
        out["标准价(*)"] = j
        out["标费档次(*)"] = l if l is not None else ""
    else:
        out["标费档次(*)"] = ""
        out["标准价(*)"] = ""
        warnings.append("款式「%s」未匹配到【新品类层级编码及标费档次】H列（中类），标准价/标费档次留空。" % style_raw)
    cost = parse_cost(row.get("成本", ""))
    if cost is None:
        cost_price = ""
        warnings.append("款号「%s」成本「%s」非数字，成本价/吊牌价留空。" % (款号, row.get("成本")))
    else:
        cost_price = cost + 10
    out["成本价(*)"] = cost_price
    if isinstance(cost_price, (int, float)):
        first_digit = str(int(cost_price))[0]
        out["吊牌价(*)"] = first_digit + "99"
    else:
        out["吊牌价(*)"] = ""
    out["生产厂商(*)"] = st["生产厂商"]
    out["渠道(*)"] = st["渠道"]
    out["平台(*)"] = st["平台"]
    out["商品备注"] = ""
    out["辅料(*)"] = st["辅料"]
    # 返回：行数据列表 + 是否需要标黄标记（新增标记）
    row_data = [out.get(h, "") for h in OUTPUT_HEADERS]
    # 标黄判断仍基于原始面料文本，避免标准化后丢失“里料/填充物/底料”等提示词。
    need_yellow = bool(re.search(r"里料|填充物|底料", fabric_raw))
    return row_data, need_yellow
# ============================================================
# 5. 读取输入表 + 生成输出
# ============================================================
INPUT_COLUMNS = ["款号", "款式", "尺码", "颜色", "上下装", "面料", "成本"]
def find_header_row(ws):
    for r in range(1, min(ws.max_row, 20) + 1):
        text = ""
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if v is not None:
                text += str(v) + "|"
        if "款号" in text and "款式" in text:
            return r
    return None
def build_column_map(ws, header_row):
    col_map = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(header_row, c).value
        if v is not None:
            key = normalize_key(v)
            if key and key not in col_map:
                col_map[key] = c
    return col_map
def generate(input_path, master_path, output_path, settings, log_cb=None):
    def log(msg):
        if log_cb:
            log_cb(msg)
    log("正在读取输入表：%s" % input_path)
    wbin = openpyxl.load_workbook(input_path, data_only=True)
    ws_in = wbin.worksheets[0]
    header_row = find_header_row(ws_in)
    if header_row is None:
        raise ValueError("输入表未找到表头行（需包含【款号】【款式】列）。")
    col_map = build_column_map(ws_in, header_row)
    col_idx = {}
    for name in INPUT_COLUMNS:
        col_idx[name] = col_map.get(normalize_key(name))
    missing_cols = [n for n in INPUT_COLUMNS if col_idx[n] is None]
    if missing_cols:
        raise ValueError("输入表缺少列：%s" % "、".join(missing_cols))
    log("主数据表：%s" % master_path)
    cat_map, price_map, size_map, color_map = load_master(master_path)
    missing_colors = []
    color_issues = []
    cfg = {
        "maps": (cat_map, price_map, size_map, color_map),
        "settings": settings,
        "missing_colors": missing_colors,
        "color_issues": color_issues,
    }
    warnings = []
    rows_out = []
    yellow_flags = []  # 新增：每一行是否面料成分要标黄
    data_rows = list(range(header_row + 1, ws_in.max_row + 1))
    total = len([r for r in data_rows if any(ws_in.cell(r, c).value is not None for c in range(1, ws_in.max_column + 1))])
    done = 0
    for r in data_rows:
        row_vals = {}
        for name in INPUT_COLUMNS:
            idx = col_idx[name]
            row_vals[name] = ws_in.cell(r, idx).value if idx else None
        if all(v is None or str(v).strip() == "" for v in row_vals.values()):
            continue
        row_data, need_yellow = process_row(row_vals, cfg, warnings)
        rows_out.append(row_data)
        yellow_flags.append(need_yellow)
        done += 1
        if log_cb and (done % 50 == 0 or done == total):
            log("已处理 %d/%d 行..." % (done, total))
    log("共生成 %d 行数据。" % len(rows_out))
    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = "Worksheet"
    ws_out.append(OUTPUT_HEADERS)
    # 表头行套用示例模板格式：加粗、灰色底纹、居中、细边框
    for c in range(1, len(OUTPUT_HEADERS) + 1):
        hcell = ws_out.cell(row=1, column=c)
        hcell.fill = HEADER_FILL
        hcell.font = HEADER_FONT
        hcell.alignment = HEADER_ALIGN
        hcell.border = HEADER_BORDER
    # 找到面料成分(*)所在列索引（从1开始）
    fabric_col_idx = OUTPUT_HEADERS.index("面料成分(*)") + 1
    for row_idx, (row_data, flag_yellow) in enumerate(zip(rows_out, yellow_flags), start=2):
        ws_out.append(row_data)
        if flag_yellow:
            cell = ws_out.cell(row=row_idx, column=fabric_col_idx)
            cell.fill = YELLOW_FILL

    for idx, _ in enumerate(OUTPUT_HEADERS, start=1):
        ws_out.column_dimensions[openpyxl.utils.get_column_letter(idx)].width = 14
    wb_out.save(output_path)
    log("输出完成：%s" % output_path)
    return warnings, len(rows_out), missing_colors, color_issues
def append_missing_colors(master_path, selections, log_cb=None):
    def log(msg):
        if log_cb:
            log_cb(msg)
    if not master_path or not os.path.exists(master_path):
        raise FileNotFoundError("找不到主数据表：%s" % master_path)
    wb = openpyxl.load_workbook(master_path)
    if "颜色" not in wb.sheetnames:
        raise KeyError("主数据表缺少【颜色】工作表")
    ws = wb["颜色"]
    existing = set()
    for r in range(2, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if a is not None and str(a).strip():
            existing.add(str(a).strip())
    added = 0
    next_row = ws.max_row + 1
    for orig, sel_name, sel_code in selections:
        if not sel_code:
            continue
        if orig in existing:
            log("颜色「%s」已存在于颜色表，跳过。" % orig)
            continue
        ws.cell(next_row, 1).value = orig
        ws.cell(next_row, 2).value = sel_code
        existing.add(orig)
        added += 1
        log("已新增颜色：「%s」-> 编码「%s」（取自所选「%s」）。" % (orig, sel_code, sel_name))
        next_row += 1
    if added:
        wb.save(master_path)
        log("已将 %d 条缺失颜色追加到【颜色】工作表末尾。" % added)
    else:
        log("没有需要新增的颜色。")
    return added
# ============================================================
# 6. 现代化界面（仅重做 UI；匹配/生成逻辑保持不变）
# ============================================================
UI = {
    "bg": "#F4F7FB",
    "card": "#FFFFFF",
    "text": "#172033",
    "muted": "#6B7280",
    "line": "#E5EAF2",
    "accent": "#2563EB",
    "accent_hover": "#1D4ED8",
    "accent_soft": "#EFF6FF",
    "success": "#059669",
    "warning": "#D97706",
    "danger": "#DC2626",
    "log_bg": "#F8FAFC",
    "log_text": "#334155",
}


def apply_modern_theme(root):
    """纯 Tkinter/ttk 现代化浅色主题，不引入第三方 UI 依赖。"""
    style = ttk.Style(root)
    try:
        style.theme_use("clam")
    except tk.TclError:
        pass

    default_font = ("Microsoft YaHei UI", 10)
    small_font = ("Microsoft YaHei UI", 9)
    bold_font = ("Microsoft YaHei UI", 10, "bold")

    root.configure(bg=UI["bg"])

    style.configure(".", font=default_font, background=UI["bg"], foreground=UI["text"])
    style.configure("App.TFrame", background=UI["bg"])
    style.configure("Card.TFrame", background=UI["card"])
    style.configure("Soft.TFrame", background=UI["accent_soft"])

    style.configure("TLabel", background=UI["bg"], foreground=UI["text"])
    style.configure("Card.TLabel", background=UI["card"], foreground=UI["text"])
    style.configure("Muted.TLabel", background=UI["card"], foreground=UI["muted"], font=small_font)
    style.configure("Section.TLabel", background=UI["card"], foreground=UI["text"],
                    font=("Microsoft YaHei UI", 11, "bold"))
    style.configure("Field.TLabel", background=UI["card"], foreground=UI["muted"], font=small_font)
    style.configure("Status.TLabel", background=UI["accent_soft"], foreground=UI["accent"], font=bold_font)

    style.configure(
        "Modern.TEntry",
        fieldbackground="#FFFFFF",
        foreground=UI["text"],
        bordercolor="#CBD5E1",
        lightcolor="#CBD5E1",
        darkcolor="#CBD5E1",
        insertcolor=UI["text"],
        padding=(9, 5),
        relief="flat",
    )
    style.map(
        "Modern.TEntry",
        bordercolor=[("focus", UI["accent"])],
        lightcolor=[("focus", UI["accent"])],
        darkcolor=[("focus", UI["accent"])],
    )

    style.configure(
        "Modern.TCombobox",
        fieldbackground="#FFFFFF",
        background="#FFFFFF",
        foreground=UI["text"],
        arrowcolor=UI["muted"],
        bordercolor="#CBD5E1",
        lightcolor="#CBD5E1",
        darkcolor="#CBD5E1",
        padding=(8, 6),
    )
    style.map(
        "Modern.TCombobox",
        bordercolor=[("focus", UI["accent"])],
        lightcolor=[("focus", UI["accent"])],
        darkcolor=[("focus", UI["accent"])],
        fieldbackground=[("readonly", "#FFFFFF")],
        foreground=[("readonly", UI["text"])],
    )

    style.configure(
        "Accent.TButton",
        background=UI["accent"],
        foreground="#FFFFFF",
        font=("Microsoft YaHei UI", 10, "bold"),
        padding=(18, 10),
        relief="flat",
        borderwidth=0,
    )
    style.map(
        "Accent.TButton",
        background=[("active", UI["accent_hover"]), ("pressed", UI["accent_hover"]), ("disabled", "#94A3B8")],
        foreground=[("disabled", "#F8FAFC")],
    )

    style.configure(
        "Secondary.TButton",
        background="#FFFFFF",
        foreground=UI["text"],
        padding=(12, 5),
        relief="flat",
        borderwidth=1,
        bordercolor="#D8E0EA",
    )
    style.map(
        "Secondary.TButton",
        background=[("active", "#F8FAFC"), ("pressed", "#F1F5F9")],
        bordercolor=[("active", "#B8C4D4")],
    )

    style.configure(
        "Danger.TButton",
        background="#FFFFFF",
        foreground=UI["danger"],
        padding=(12, 8),
        relief="flat",
        borderwidth=1,
        bordercolor="#F3C7C7",
    )
    style.map("Danger.TButton", background=[("active", "#FEF2F2")])

    style.configure(
        "Modern.Horizontal.TProgressbar",
        troughcolor="#E8EEF7",
        background=UI["accent"],
        bordercolor="#E8EEF7",
        lightcolor=UI["accent"],
        darkcolor=UI["accent"],
        thickness=5,
    )

    style.configure("Modern.Vertical.TScrollbar", background="#D9E1EC", troughcolor=UI["bg"],
                    arrowcolor=UI["muted"], borderwidth=0)
    return style


class MissingColorRow(ttk.Frame):
    def __init__(self, parent, orig, sugs, full_color_list):
        super().__init__(parent, style="Card.TFrame")
        self.orig = orig
        self.sugs = sugs
        self.full_color_list = full_color_list  # 全部颜色库
        self.selected = None
        self.search_var = tk.StringVar()

        self.columnconfigure(2, weight=1)

        orig_wrap = ttk.Frame(self, style="Soft.TFrame", padding=(10, 7))
        orig_wrap.grid(row=0, column=0, sticky="ew", padx=(0, 10), pady=4)
        tk.Label(
            orig_wrap, text=orig, bg=UI["accent_soft"], fg=UI["accent"],
            font=("Microsoft YaHei UI", 9, "bold"), anchor="w", width=14
        ).pack(fill="x")

        self.search_entry = ttk.Entry(
            self, textvariable=self.search_var, width=16, style="Modern.TEntry"
        )
        self.search_entry.grid(row=0, column=1, sticky="ew", padx=(0, 10), pady=4)
        self.search_var.trace_add("write", self._on_search_change)

        if sugs:
            self.selected = sugs[0]
        self.combo = ttk.Combobox(
            self, state="readonly", width=35, style="Modern.TCombobox"
        )
        self.combo.grid(row=0, column=2, sticky="ew", pady=4)
        self.combo.bind("<<ComboboxSelected>>", self._on_select)

        # 初始化下拉列表（逻辑保持原样）
        self._refresh_combo("")

    def _refresh_combo(self, keyword):
        """根据搜索词刷新下拉候选项，优先原有推荐，再追加匹配搜索的颜色"""
        kw = keyword.lower().strip()
        base_options = []
        # 原有推荐优先放前面
        if self.sugs:
            base_options.extend([("%s (%s)" % (n, c), n, c) for n, c in self.sugs])
        # 搜索过滤全部颜色库，匹配关键字
        matched = []
        for name, code in self.full_color_list:
            if kw in name.lower():
                matched.append((f"{name} ({code})", name, code))
        # 去重
        all_items = []
        seen_names = set()
        for disp, n, c in base_options + matched:
            if n not in seen_names:
                seen_names.add(n)
                all_items.append((disp, n, c))
        if not all_items:
            all_items = [("（无匹配颜色）", "", "")]
        self.combo_values_raw = all_items
        disp_list = [x[0] for x in all_items]
        self.combo["values"] = disp_list
        self.combo.current(0)
        if all_items[0][1]:
            self.selected = (all_items[0][1], all_items[0][2])
        else:
            self.selected = None

    def _on_search_change(self, *args):
        """输入框文字变化触发刷新下拉"""
        key = self.search_var.get()
        self._refresh_combo(key)

    def _on_select(self, event=None):
        idx = self.combo.current()
        if 0 <= idx < len(self.combo_values_raw):
            _, name, code = self.combo_values_raw[idx]
            if name and code:
                self.selected = (name, code)
            else:
                self.selected = None


class ColorPanel(tk.Toplevel):
    def __init__(self, master, missing_colors, master_path, log_cb=None, on_updated=None):
        super().__init__(master)
        self.master = master
        self.master_path = master_path
        self.log_cb = log_cb
        self.on_updated = on_updated
        self.rows = []

        self.title("缺失颜色处理")
        self.configure(bg=UI["bg"])
        sw = self.winfo_screenwidth()
        sh = self.winfo_screenheight()
        panel_w = min(860, max(700, sw - 120))
        panel_h = min(620, max(460, sh - 160))
        panel_w = min(panel_w, max(660, sw - 40))
        panel_h = min(panel_h, max(420, sh - 80))
        self.minsize(min(700, panel_w), min(440, panel_h))
        self.geometry(f"{panel_w}x{panel_h}")
        self.resizable(True, True)
        self.transient(master)
        self.grab_set()

        self.full_color_list = get_full_color_list(master_path)

        # 顶部标题区
        header = tk.Frame(self, bg=UI["card"], height=92)
        header.pack(fill="x")
        header.pack_propagate(False)

        title_box = tk.Frame(header, bg=UI["card"])
        title_box.pack(side="left", fill="y", padx=24, pady=18)
        tk.Label(
            title_box, text="缺失颜色处理", bg=UI["card"], fg=UI["text"],
            font=("Microsoft YaHei UI", 15, "bold")
        ).pack(anchor="w")
        tk.Label(
            title_box,
            text="确认对应颜色后会写入主数据，并自动重新生成当前输出文件。",
            bg=UI["card"], fg=UI["muted"], font=("Microsoft YaHei UI", 9)
        ).pack(anchor="w", pady=(4, 0))

        badge = tk.Label(
            header, text=f"  {len(missing_colors)} 项待处理  ",
            bg=UI["accent_soft"], fg=UI["accent"],
            font=("Microsoft YaHei UI", 9, "bold"), padx=8, pady=6
        )
        badge.pack(side="right", padx=24)

        # 主内容卡片
        body = ttk.Frame(self, style="App.TFrame", padding=(20, 16, 20, 10))
        body.pack(fill="both", expand=True)

        card = ttk.Frame(body, style="Card.TFrame", padding=16)
        card.pack(fill="both", expand=True)

        head = ttk.Frame(card, style="Card.TFrame")
        head.pack(fill="x", pady=(0, 8))
        head.columnconfigure(2, weight=1)
        ttk.Label(head, text="未匹配颜色", style="Field.TLabel", width=17).grid(row=0, column=0, sticky="w")
        ttk.Label(head, text="搜索颜色", style="Field.TLabel", width=18).grid(row=0, column=1, sticky="w", padx=(0, 10))
        ttk.Label(head, text="选择对应主数据颜色（推荐项优先）", style="Field.TLabel").grid(row=0, column=2, sticky="w")

        ttk.Separator(card, orient="horizontal").pack(fill="x", pady=(0, 8))

        scroll_container = ttk.Frame(card, style="Card.TFrame")
        scroll_container.pack(fill="both", expand=True)

        canvas = tk.Canvas(
            scroll_container, highlightthickness=0, bd=0,
            bg=UI["card"]
        )
        sbar = ttk.Scrollbar(
            scroll_container, orient="vertical", command=canvas.yview,
            style="Modern.Vertical.TScrollbar"
        )
        inner = ttk.Frame(canvas, style="Card.TFrame")
        inner.bind("<Configure>", lambda e: canvas.configure(scrollregion=canvas.bbox("all")))
        win = canvas.create_window((0, 0), window=inner, anchor="nw")
        canvas.configure(yscrollcommand=sbar.set)
        canvas.bind("<Configure>", lambda e: canvas.itemconfigure(win, width=e.width))
        canvas.bind("<MouseWheel>", lambda e: canvas.yview_scroll(int(-1 * (e.delta / 120)), "units"))
        canvas.pack(side="left", fill="both", expand=True)
        sbar.pack(side="right", fill="y")

        for mc in missing_colors:
            row = MissingColorRow(inner, mc["orig"], mc["sugs"], self.full_color_list)
            row.pack(fill="x", pady=3)
            self.rows.append(row)

        # 固定底部操作区
        footer = ttk.Frame(self, style="Card.TFrame", padding=(20, 14))
        footer.pack(fill="x")
        ttk.Button(
            footer, text="取消", command=self.destroy, style="Secondary.TButton"
        ).pack(side="right", padx=(8, 0))
        ttk.Button(
            footer, text="确认并追加到颜色表", command=self._confirm, style="Accent.TButton"
        ).pack(side="right")

    def _confirm(self):
        unselected = [r.orig for r in self.rows if r.selected is None]
        if unselected:
            messagebox.showwarning("提示", "还有颜色未选择：\n" + "、".join(unselected), parent=self)
            return
        selections = [(r.orig, r.selected[0], r.selected[1]) for r in self.rows]
        try:
            added = append_missing_colors(self.master_path, selections, self.log_cb)
        except Exception as e:
            messagebox.showerror("出错", str(e), parent=self)
            return

        self.destroy()
        if self.on_updated:
            # 颜色表写入完成后自动重新读取主数据并覆盖生成当前输出文件。
            self.master.after(80, lambda: self.on_updated(added))
        else:
            messagebox.showinfo("完成", "已把 %d 条缺失颜色追加到颜色表末尾。" % added)

    def log(self, msg):
        if self.log_cb:
            self.log_cb(msg)


class App:
    DEFAULTS = {
        "开发前缀": "KQ",
        "产品前缀": "Dickies",
        "年份": "2026",
        "季节": "秋季",
        "大类": "男装",
        "性别": "男",
        "包材单价": "0.5",
        "生产厂商": "泉州市柏香达人贸易有限公司",
        "渠道": "线上",
        "平台": "得物;抖音",
        "辅料": "主唛;织唛;洗水唛",
    }

    def __init__(self, root):
        self.root = root
        self.style = apply_modern_theme(root)

        root.title("送审资料生成器")
        # 根据屏幕自动计算初始尺寸，避免高 DPI / 小屏幕下必须全屏才能看到底部操作区。
        sw = root.winfo_screenwidth()
        sh = root.winfo_screenheight()
        win_w = min(1080, max(760, sw - 120))
        win_h = min(720, max(560, sh - 70))
        win_w = min(win_w, max(720, sw - 40))
        win_h = min(win_h, max(520, sh - 55))
        pos_x = max(0, (sw - win_w) // 2)
        pos_y = max(0, (sh - win_h) // 3)
        root.geometry(f"{win_w}x{win_h}+{pos_x}+{pos_y}")
        root.minsize(min(820, win_w), min(560, win_h))
        root.resizable(True, True)

        self.var_input = tk.StringVar()
        self.var_master = tk.StringVar(value=r"D:\GZ\预输入.xlsx")
        self.var_output = tk.StringVar()
        self.settings_vars = {}
        self.last_output_path = ""

        self._build_header()
        self._build_main()
        self._set_status("就绪", "ready")
        self.log("就绪。请选择输入表，主数据表默认读取 D:\\GZ\\预输入.xlsx。", "info")

    def _build_header(self):
        header = tk.Frame(self.root, bg="#172033", height=80)
        header.pack(fill="x")
        header.pack_propagate(False)

        title_area = tk.Frame(header, bg="#172033")
        title_area.pack(side="left", padx=24, pady=12)

        tk.Label(
            title_area, text="送审资料生成器",
            bg="#172033", fg="#FFFFFF",
            font=("Microsoft YaHei UI", 18, "bold")
        ).pack(anchor="w")
        tk.Label(
            title_area,
            text="Excel 数据整理 · 规则匹配 · 一键生成送审导入文件",
            bg="#172033", fg="#AAB6CB",
            font=("Microsoft YaHei UI", 9)
        ).pack(anchor="w", pady=(4, 0))

        self.status_badge = tk.Label(
            header, text="  ● 就绪  ",
            bg="#22304B", fg="#93C5FD",
            font=("Microsoft YaHei UI", 9, "bold"),
            padx=10, pady=6
        )
        self.status_badge.pack(side="right", padx=24)

    def _build_main(self):
        # 主内容支持纵向滚动：窗口无需全屏，较小屏幕 / 高 DPI 下也能访问全部 UI。
        shell = tk.Frame(self.root, bg=UI["bg"])
        shell.pack(fill="both", expand=True)

        self.main_canvas = tk.Canvas(
            shell, bg=UI["bg"], highlightthickness=0, bd=0
        )
        self.main_scrollbar = ttk.Scrollbar(
            shell, orient="vertical", command=self.main_canvas.yview,
            style="Modern.Vertical.TScrollbar"
        )
        self.main_canvas.configure(yscrollcommand=self.main_scrollbar.set)
        self.main_canvas.pack(side="left", fill="both", expand=True)
        self.main_scrollbar.pack(side="right", fill="y")

        outer = ttk.Frame(self.main_canvas, style="App.TFrame", padding=(18, 12, 18, 14))
        self._main_window = self.main_canvas.create_window((0, 0), window=outer, anchor="nw")

        def sync_scrollregion(event=None):
            self.main_canvas.configure(scrollregion=self.main_canvas.bbox("all"))

        def fit_content_width(event):
            self.main_canvas.itemconfigure(self._main_window, width=event.width)

        outer.bind("<Configure>", sync_scrollregion)
        self.main_canvas.bind("<Configure>", fit_content_width)

        def on_main_mousewheel(event):
            # 日志框、颜色弹窗等拥有自己的滚动行为，不带动主窗口。
            try:
                if event.widget.winfo_toplevel() is not self.root:
                    return
            except Exception:
                return
            if isinstance(event.widget, tk.Text):
                return
            bbox = self.main_canvas.bbox("all")
            if bbox and bbox[3] > self.main_canvas.winfo_height():
                self.main_canvas.yview_scroll(int(-1 * (event.delta / 120)), "units")

        self.root.bind_all("<MouseWheel>", on_main_mousewheel, add="+")

        # 文件区
        file_card = ttk.Frame(outer, style="Card.TFrame", padding=15)
        file_card.pack(fill="x")

        title_row = ttk.Frame(file_card, style="Card.TFrame")
        title_row.pack(fill="x", pady=(0, 8))
        ttk.Label(title_row, text="文件", style="Section.TLabel").pack(side="left")
        ttk.Label(
            title_row, text="仅支持 .xlsx；输出文件会自动补全建议名称",
            style="Muted.TLabel"
        ).pack(side="right")

        files_grid = ttk.Frame(file_card, style="Card.TFrame")
        files_grid.pack(fill="x")
        files_grid.columnconfigure(1, weight=1)

        self._add_file_row(
            files_grid, 0, "输入表", "款号 / 款式 / 尺码 / 颜色 / 上下装 / 面料 / 成本",
            self.var_input, self.choose_input
        )
        self._add_file_row(
            files_grid, 1, "主数据表", "默认 D:\\GZ\\预输入.xlsx",
            self.var_master, self.choose_master
        )
        self._add_file_row(
            files_grid, 2, "输出位置", "生成后的送审导入 Excel",
            self.var_output, self.choose_output
        )

        # 中间双列：设置 + 日志
        middle = ttk.Frame(outer, style="App.TFrame")
        middle.pack(fill="both", expand=True, pady=(10, 0))
        middle.columnconfigure(0, weight=3, uniform="main")
        middle.columnconfigure(1, weight=2, uniform="main")
        middle.rowconfigure(0, weight=1)

        self._build_settings_card(middle)
        self._build_log_card(middle)

        # 底部操作栏
        actions = ttk.Frame(outer, style="App.TFrame")
        actions.pack(fill="x", pady=(10, 0))

        self.progress = ttk.Progressbar(
            actions, mode="indeterminate", style="Modern.Horizontal.TProgressbar"
        )
        self.progress.pack(side="left", fill="x", expand=True, padx=(0, 14))

        self.open_btn = ttk.Button(
            actions, text="打开输出文件夹", command=self.open_output_folder,
            style="Secondary.TButton", state="disabled"
        )
        self.open_btn.pack(side="right", padx=(8, 0))

        self.btn = ttk.Button(
            actions, text="开始生成", command=self.run, style="Accent.TButton"
        )
        self.btn.pack(side="right")

    def _build_settings_card(self, parent):
        card = ttk.Frame(parent, style="Card.TFrame", padding=15)
        card.grid(row=0, column=0, sticky="nsew", padx=(0, 7))

        head = ttk.Frame(card, style="Card.TFrame")
        head.pack(fill="x", pady=(0, 8))
        ttk.Label(head, text="固定项设置", style="Section.TLabel").pack(side="left")
        ttk.Button(
            head, text="恢复默认", command=self.reset_defaults, style="Secondary.TButton"
        ).pack(side="right")

        ttk.Label(
            card, text="这些字段会直接写入生成文件，可按本批次调整。",
            style="Muted.TLabel"
        ).pack(anchor="w", pady=(0, 8))

        grid = ttk.Frame(card, style="Card.TFrame")
        grid.pack(fill="both", expand=True)
        grid.columnconfigure(0, weight=1)
        grid.columnconfigure(1, weight=1)
        grid.columnconfigure(2, weight=1)

        keys = list(self.DEFAULTS.keys())
        for i, key in enumerate(keys):
            # 三列布局降低整体高度，默认窗口即可完整显示更多内容。
            r = i // 3
            c = i % 3
            field = ttk.Frame(grid, style="Card.TFrame")
            field.grid(
                row=r, column=c, sticky="ew",
                padx=(0, 7) if c == 0 else ((7, 7) if c == 1 else (7, 0)), pady=(0, 4)
            )
            ttk.Label(field, text=key, style="Field.TLabel").pack(anchor="w", pady=(0, 4))
            v = tk.StringVar(value=self.DEFAULTS[key])
            ttk.Entry(
                field, textvariable=v, style="Modern.TEntry"
            ).pack(fill="x")
            self.settings_vars[key] = v

    def _build_log_card(self, parent):
        card = ttk.Frame(parent, style="Card.TFrame", padding=15)
        card.grid(row=0, column=1, sticky="nsew", padx=(7, 0))

        head = ttk.Frame(card, style="Card.TFrame")
        head.pack(fill="x", pady=(0, 8))
        ttk.Label(head, text="运行日志", style="Section.TLabel").pack(side="left")
        ttk.Button(
            head, text="清空", command=self.clear_log, style="Secondary.TButton"
        ).pack(side="right")

        ttk.Label(
            card, text="未匹配、格式问题和生成状态会显示在这里。",
            style="Muted.TLabel"
        ).pack(anchor="w", pady=(0, 8))

        log_wrap = tk.Frame(card, bg=UI["log_bg"], highlightthickness=1,
                            highlightbackground=UI["line"])
        log_wrap.pack(fill="both", expand=True)

        self.txt = tk.Text(
            log_wrap, wrap="word", height=8, font=("Microsoft YaHei UI", 9),
            bg=UI["log_bg"], fg=UI["log_text"],
            selectbackground="#DBEAFE", selectforeground=UI["text"],
            relief="flat", bd=0, padx=12, pady=12,
            insertbackground=UI["text"]
        )
        sb = ttk.Scrollbar(
            log_wrap, command=self.txt.yview, style="Modern.Vertical.TScrollbar"
        )
        self.txt.configure(yscrollcommand=sb.set)
        sb.pack(side="right", fill="y")
        self.txt.pack(side="left", fill="both", expand=True)

        self.txt.tag_configure("info", foreground=UI["log_text"])
        self.txt.tag_configure("success", foreground=UI["success"])
        self.txt.tag_configure("warning", foreground=UI["warning"])
        self.txt.tag_configure("error", foreground=UI["danger"])
        self.txt.configure(state="disabled")

    def _add_file_row(self, parent, r, label, hint, var, cmd):
        label_box = ttk.Frame(parent, style="Card.TFrame")
        label_box.grid(row=r, column=0, sticky="w", padx=(0, 12), pady=2)
        # 标签与说明改为单行，降低文件区高度；窄屏仍可通过主界面滚动查看全部内容。
        ttk.Label(
            label_box, text=label, style="Card.TLabel",
            font=("Microsoft YaHei UI", 9, "bold")
        ).pack(side="left")
        ttk.Label(
            label_box, text=" · " + hint, style="Muted.TLabel"
        ).pack(side="left", padx=(4, 0))

        ttk.Entry(
            parent, textvariable=var, style="Modern.TEntry"
        ).grid(row=r, column=1, sticky="ew", pady=2)

        ttk.Button(
            parent, text="选择文件", command=cmd, style="Secondary.TButton"
        ).grid(row=r, column=2, padx=(10, 0), pady=2)

    def choose_input(self):
        p = filedialog.askopenfilename(
            title="选择输入表",
            filetypes=[("Excel 工作簿", "*.xlsx"), ("所有文件", "*.*")]
        )
        if p:
            self.var_input.set(p)
            if not self.var_output.get():
                base = os.path.splitext(p)[0]
                self.var_output.set(base + "_送审导入.xlsx")

    def choose_master(self):
        p = filedialog.askopenfilename(
            title="选择主数据表",
            filetypes=[("Excel 工作簿", "*.xlsx"), ("所有文件", "*.*")]
        )
        if p:
            self.var_master.set(p)

    def choose_output(self):
        p = filedialog.asksaveasfilename(
            title="保存输出到",
            defaultextension=".xlsx",
            filetypes=[("Excel 工作簿", "*.xlsx")]
        )
        if p:
            self.var_output.set(p)

    def clear_log(self):
        self.txt.configure(state="normal")
        self.txt.delete("1.0", "end")
        self.txt.configure(state="disabled")

    def reset_defaults(self):
        for key, value in self.DEFAULTS.items():
            if key in self.settings_vars:
                self.settings_vars[key].set(value)
        self.log("已恢复固定项默认值。", "info")

    def _set_status(self, text, kind="ready"):
        palette = {
            "ready": ("#22304B", "#93C5FD"),
            "running": ("#293753", "#FDE68A"),
            "success": ("#183E37", "#6EE7B7"),
            "error": ("#492A34", "#FCA5A5"),
        }
        bg, fg = palette.get(kind, palette["ready"])
        self.status_badge.configure(text=f"  ● {text}  ", bg=bg, fg=fg)

    def log(self, msg, level="info"):
        self.txt.configure(state="normal")
        self.txt.insert("end", msg + "\n", level)
        self.txt.see("end")
        self.txt.configure(state="disabled")
        self.root.update_idletasks()

    def get_settings(self):
        return {k: v.get().strip() for k, v in self.settings_vars.items()}

    def open_color_panel(self, missing_colors, master_path):
        if not missing_colors:
            return
        seen = set()
        dedup = []
        for mc in missing_colors:
            o = mc["orig"]
            if o not in seen:
                seen.add(o)
                dedup.append(mc)
        ColorPanel(
            self.root, dedup, master_path, self.log,
            on_updated=self._regenerate_after_color_update
        )

    def _regenerate_after_color_update(self, added):
        self.log("─" * 46, "info")
        self.log("颜色表已更新 %d 条，正在自动重新生成输出文件..." % added, "success")
        self.run(auto_reason="color_update")

    def _color_issue_summary(self, color_issues, limit=10):
        lines = []
        for issue in color_issues[:limit]:
            reasons = []
            if issue.get("缺失颜色"):
                reasons.append("未匹配：" + "、".join(issue["缺失颜色"]))
            dup = issue.get("重复代码") or {}
            if dup:
                reasons.append("重复编码：" + "；".join(
                    "%s=%s" % (code, "、".join(cols)) for code, cols in dup.items()
                ))
            if not reasons:
                reasons.append("颜色数与编码数不一致")
            lines.append(
                "款号 %s｜%s｜%d色→%d个唯一代码｜%s"
                % (issue.get("款号") or "未填写", issue.get("款式") or "未填写",
                   issue.get("原颜色数", 0), issue.get("唯一代码数", 0), "；".join(reasons))
            )
        if len(color_issues) > limit:
            lines.append("……另有 %d 款，请查看右侧运行日志。" % (len(color_issues) - limit))
        return "\n".join(lines)

    def _validate_paths(self, input_path, master_path, output_path):
        """只做文件层校验，不改变任何业务匹配逻辑。"""
        if not input_path:
            return "请先选择输入表。"
        if not master_path:
            return "请先选择主数据表。"
        if not input_path.lower().endswith(".xlsx"):
            return "输入表需为 .xlsx 格式（openpyxl 不支持旧版 .xls）。"
        if not master_path.lower().endswith(".xlsx"):
            return "主数据表需为 .xlsx 格式（openpyxl 不支持旧版 .xls）。"
        if output_path and not output_path.lower().endswith(".xlsx"):
            return "输出文件需保存为 .xlsx 格式。"
        return ""

    def open_output_folder(self):
        path = self.last_output_path or self.var_output.get().strip()
        if not path:
            return
        folder = os.path.dirname(os.path.abspath(path))
        try:
            if os.name == "nt":
                os.startfile(folder)
            else:
                messagebox.showinfo("输出目录", folder)
        except Exception as e:
            messagebox.showerror("无法打开目录", str(e))

    def run(self, auto_reason=None):
        input_path = self.var_input.get().strip()
        master_path = self.var_master.get().strip()
        output_path = self.var_output.get().strip()

        if not output_path and input_path:
            base = os.path.splitext(input_path)[0]
            output_path = base + "_送审导入.xlsx"
            self.var_output.set(output_path)

        path_error = self._validate_paths(input_path, master_path, output_path)
        if path_error:
            messagebox.showwarning("请检查文件", path_error)
            self._set_status("待检查", "error")
            return

        settings = self.get_settings()
        if auto_reason != "color_update":
            self.clear_log()
            self.log("开始生成...", "info")
        else:
            self.log("重新读取更新后的颜色表...", "info")
        self._set_status("生成中", "running")
        self.btn.config(state="disabled")
        self.open_btn.config(state="disabled")
        self.progress.start(12)

        try:
            warnings, count, missing_colors, color_issues = generate(
                input_path, master_path, output_path, settings, self.log
            )
            self.log("─" * 46, "info")
            if warnings:
                self.log("存在 %d 条提示，请核对：" % len(warnings), "warning")
                for w in warnings:
                    self.log("• " + w, "warning")
            else:
                self.log("全部字段生成无提示。", "success")

            if color_issues:
                self.log("颜色完整性检测发现 %d 款异常：" % len(color_issues), "warning")
                for issue in color_issues:
                    self.log(
                        "• 款号「%s」款式「%s」｜原始颜色：%s｜原%d色 / 匹配%d码 / 唯一%d码"
                        % (issue.get("款号") or "未填写", issue.get("款式") or "未填写",
                           issue.get("原始颜色") or "", issue.get("原颜色数", 0),
                           issue.get("匹配代码数", 0), issue.get("唯一代码数", 0)),
                        "warning"
                    )

            self.last_output_path = output_path
            self.open_btn.config(state="normal")
            self._set_status("生成完成", "success")

            if missing_colors:
                messagebox.showwarning(
                    "生成完成｜存在缺失颜色",
                    "已生成 %d 行数据，但存在未匹配颜色。\n\n"
                    "请在接下来的窗口确认对应颜色；确认后程序会自动重新生成并覆盖当前输出文件。"
                    % count
                )
                self.open_color_panel(missing_colors, master_path)
            elif color_issues:
                messagebox.showwarning(
                    "生成完成｜颜色检测异常",
                    "输出文件已生成，但颜色检测发现以下款式需要核对：\n\n%s\n\n完整信息已写入运行日志。"
                    % self._color_issue_summary(color_issues)
                )
            else:
                messagebox.showinfo(
                    "生成完成",
                    "已成功生成 %d 行数据。\n\n保存位置：\n%s" % (count, output_path)
                )

        except Exception as e:
            self.log("出错：" + str(e), "error")
            self._set_status("生成失败", "error")
            messagebox.showerror("生成失败", str(e))
        finally:
            self.progress.stop()
            self.btn.config(state="normal")


def main():
    root = tk.Tk()
    App(root)
    root.mainloop()


if __name__ == "__main__":
    main()
