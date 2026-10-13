import os
import re
import traceback
from collections import defaultdict
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

import pandas as pd
from openpyxl import Workbook, load_workbook


class ExcelMatcherApp:
    """聚水潭 / 抖店编码匹配工具。

    核心关系：
    - 一个聚水潭款式可以绑定多个抖店 PID；
    - 同一个抖店 PID 也可以同时绑定多个聚水潭款式；
    - 当同一个 PID 绑定多个聚水潭款式时，依靠“抖店颜色 -> 聚水潭款式/颜色”唯一归属，
      再结合尺码定位最终聚水潭 SKU 商品编码。
    """

    BG = "#F5F7FA"
    CARD = "#FFFFFF"
    TEXT = "#303133"
    MUTED = "#909399"
    BORDER = "#E4E7ED"
    PRIMARY = "#409EFF"
    SUCCESS = "#67C23A"
    WARNING = "#E6A23C"
    DANGER = "#F56C6C"

    def __init__(self, root):
        self.root = root
        self.root.title("聚水潭 - 抖店编码匹配工具")
        self.root.geometry("1120x780")
        self.root.minsize(980, 680)
        self.root.configure(bg=self.BG)

        self._setup_styles()
        self._init_data()
        self._build_shell()
        self.show_file_page()

    # ----------------------------- 基础工具 -----------------------------
    @staticmethod
    def norm_id(value):
        """统一 Excel 编码的显示形式，尽量避免 123.0 / 空值造成匹配失败。"""
        if value is None:
            return ""
        try:
            if pd.isna(value):
                return ""
        except Exception:
            pass
        if isinstance(value, bool):
            return str(value)
        if isinstance(value, int):
            return str(value)
        if isinstance(value, float):
            if value.is_integer():
                return str(int(value))
            return format(value, ".15g")
        text = str(value).strip()
        if re.fullmatch(r"[-+]?\d+\.0+", text):
            return text.split(".", 1)[0]
        return text

    @staticmethod
    def norm_text(value):
        if value is None:
            return ""
        try:
            if pd.isna(value):
                return ""
        except Exception:
            pass
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value).strip()

    def _setup_styles(self):
        self.style = ttk.Style()
        try:
            self.style.theme_use("clam")
        except tk.TclError:
            pass

        font = "微软雅黑"
        self.style.configure("App.TFrame", background=self.BG)
        self.style.configure("Card.TFrame", background=self.CARD)
        self.style.configure("Title.TLabel", font=(font, 16, "bold"), foreground="#1F2D3D", background=self.BG)
        self.style.configure("Step.TLabel", font=(font, 10, "bold"), foreground=self.PRIMARY, background=self.BG)
        self.style.configure("SubTitle.TLabel", font=(font, 11, "bold"), foreground=self.TEXT, background=self.CARD)
        self.style.configure("Body.TLabel", font=(font, 10), foreground="#606266", background=self.BG)
        self.style.configure("CardBody.TLabel", font=(font, 10), foreground="#606266", background=self.CARD)
        self.style.configure("Muted.TLabel", font=(font, 9), foreground=self.MUTED, background=self.BG)
        self.style.configure("Primary.TButton", font=(font, 10, "bold"), padding=(16, 9))
        self.style.configure("Ghost.TButton", font=(font, 10), padding=(12, 8))
        self.style.configure("Small.TButton", font=(font, 9), padding=(8, 5))
        self.style.configure("TCombobox", padding=5)
        self.style.configure("TLabelframe", background=self.CARD, bordercolor=self.BORDER, relief="solid")
        self.style.configure("TLabelframe.Label", font=(font, 10, "bold"), foreground=self.TEXT, background=self.CARD)

    def _init_data(self):
        self.jushuitan_path = ""
        self.doudian_path = ""
        self.jushuitan_df = None
        self.doudian_df = None
        self.js_header_idx = 0
        self.dd_header_idx = 1

        self.style_list = []
        self.dd_pid_list = []
        self.style_pid_rows = {}      # style -> [row_info]
        self.style_to_pid = {}        # style -> [pid]
        self.style_card_map = {}      # style -> {card, rows_frame, ...}

        self.color_mapping = {}       # style -> pid -> {js_color: dd_color}
        self.temp_selected = {}       # (style, pid) -> js_color
        self.js_listbox_map = {}
        self.dd_listbox_map = {}
        self.pair_label_map = {}
        self.block_meta = {}          # (style,pid) -> js_colors/dd_colors

        self.page = None
        self._mousewheel_canvas = None

    def _build_shell(self):
        top = tk.Frame(self.root, bg="#FFFFFF", height=74, highlightbackground=self.BORDER, highlightthickness=1)
        top.pack(fill="x")
        top.pack_propagate(False)
        title_wrap = tk.Frame(top, bg="#FFFFFF")
        title_wrap.pack(fill="both", expand=True, padx=28, pady=13)
        tk.Label(title_wrap, text="编码匹配工具", font=("微软雅黑", 17, "bold"), fg="#1F2D3D", bg="#FFFFFF").pack(anchor="w")
        self.shell_subtitle = tk.Label(title_wrap, text="聚水潭 SKU 编码 → 抖店商家编码", font=("微软雅黑", 9), fg=self.MUTED, bg="#FFFFFF")
        self.shell_subtitle.pack(anchor="w", pady=(2, 0))

        self.container = tk.Frame(self.root, bg=self.BG)
        self.container.pack(fill="both", expand=True, padx=28, pady=22)

    def _clear_page(self):
        if self.page is not None:
            self.page.destroy()
        self.page = tk.Frame(self.container, bg=self.BG)
        self.page.pack(fill="both", expand=True)

    def _page_header(self, step_text, title, desc):
        ttk.Label(self.page, text=step_text, style="Step.TLabel").pack(anchor="w")
        ttk.Label(self.page, text=title, style="Title.TLabel").pack(anchor="w", pady=(4, 4))
        ttk.Label(self.page, text=desc, style="Body.TLabel", wraplength=1000).pack(anchor="w", pady=(0, 16))

    def _make_scroller(self, parent):
        wrap = tk.Frame(parent, bg=self.BG)
        canvas = tk.Canvas(wrap, bg=self.BG, highlightthickness=0)
        scrollbar = ttk.Scrollbar(wrap, orient="vertical", command=canvas.yview)
        inner = tk.Frame(canvas, bg=self.BG)
        window_id = canvas.create_window((0, 0), window=inner, anchor="nw")

        def sync_scrollregion(_event=None):
            canvas.configure(scrollregion=canvas.bbox("all"))

        def sync_width(event):
            canvas.itemconfigure(window_id, width=event.width)

        inner.bind("<Configure>", sync_scrollregion)
        canvas.bind("<Configure>", sync_width)
        canvas.configure(yscrollcommand=scrollbar.set)
        canvas.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")
        self._bind_mousewheel(canvas)
        return wrap, canvas, inner

    def _bind_mousewheel(self, canvas):
        self._mousewheel_canvas = canvas

        def _wheel(event):
            if self._mousewheel_canvas is None:
                return
            if getattr(event, "num", None) == 4:
                self._mousewheel_canvas.yview_scroll(-1, "units")
            elif getattr(event, "num", None) == 5:
                self._mousewheel_canvas.yview_scroll(1, "units")
            elif getattr(event, "delta", 0):
                delta = int(-event.delta / 120) or (-1 if event.delta > 0 else 1)
                self._mousewheel_canvas.yview_scroll(delta, "units")

        # 只绑定一次到 root，避免多页面 bind_all 叠加导致滚动越来越快。
        if not getattr(self, "_wheel_bound", False):
            self.root.bind_all("<MouseWheel>", _wheel)
            self.root.bind_all("<Button-4>", _wheel)
            self.root.bind_all("<Button-5>", _wheel)
            self._wheel_bound = True

    def _detect_header_row(self, path, required_columns, max_scan=12):
        raw = pd.read_excel(path, header=None, nrows=max_scan, dtype=object)
        required = {str(c).strip() for c in required_columns}
        for idx, row in raw.iterrows():
            values = {self.norm_text(v) for v in row.tolist() if self.norm_text(v)}
            if required.issubset(values):
                return int(idx)
        raise ValueError(f"前 {max_scan} 行中未找到表头：{', '.join(required_columns)}")

    # ----------------------------- 第一步 -----------------------------
    def show_file_page(self):
        self._clear_page()
        self._page_header("步骤 1 / 3", "选择数据文件", "先选择聚水潭编码表和抖店商品表。程序会自动识别表头位置并检查必要字段。")

        card = tk.Frame(self.page, bg=self.CARD, highlightbackground=self.BORDER, highlightthickness=1)
        card.pack(fill="x", pady=(4, 8), ipady=8)

        self.js_file_name = self._file_picker_row(card, "聚水潭编码文件", "包含：款式编码、颜色、规格、商品编码", self.choose_js_file)
        self.dd_file_name = self._file_picker_row(card, "抖店编码文件", "包含：商品id(必填)、商品规格、商家编码(必填)", self.choose_dd_file)

        if self.jushuitan_path:
            self.js_file_name.config(text=os.path.basename(self.jushuitan_path), fg=self.TEXT)
        if self.doudian_path:
            self.dd_file_name.config(text=os.path.basename(self.doudian_path), fg=self.TEXT)

        tip = tk.Frame(self.page, bg="#ECF5FF", highlightbackground="#B3D8FF", highlightthickness=1)
        tip.pack(fill="x", pady=12)
        tk.Label(
            tip,
            text="兼容场景：一个聚水潭款式 ↔ 多个抖店 PID；同一个抖店 PID ↔ 多个聚水潭款式。后者会按颜色 + 尺码自动定位正确 SKU。",
            bg="#ECF5FF", fg="#337ECC", font=("微软雅黑", 9), anchor="w", justify="left", wraplength=1000
        ).pack(fill="x", padx=14, pady=10)

        bottom = tk.Frame(self.page, bg=self.BG)
        bottom.pack(fill="x", pady=20)
        ttk.Button(bottom, text="下一步：款式与 PID 配对", style="Primary.TButton", command=self.go_style_page).pack(side="right")

    def _file_picker_row(self, parent, title, hint, command):
        row = tk.Frame(parent, bg=self.CARD)
        row.pack(fill="x", padx=18, pady=12)
        left = tk.Frame(row, bg=self.CARD)
        left.pack(side="left", fill="x", expand=True)
        tk.Label(left, text=title, bg=self.CARD, fg=self.TEXT, font=("微软雅黑", 10, "bold")).pack(anchor="w")
        name = tk.Label(left, text=hint, bg=self.CARD, fg=self.MUTED, font=("微软雅黑", 9), anchor="w")
        name.pack(anchor="w", pady=(4, 0))
        ttk.Button(row, text="选择文件", command=command, style="Ghost.TButton", width=12).pack(side="right")
        return name

    def choose_js_file(self):
        path = filedialog.askopenfilename(filetypes=[("Excel 文件", "*.xlsx *.xls")])
        if path:
            self.jushuitan_path = path
            self.js_file_name.config(text=os.path.basename(path), fg=self.TEXT)

    def choose_dd_file(self):
        path = filedialog.askopenfilename(filetypes=[("Excel 文件", "*.xlsx *.xls")])
        if path:
            self.doudian_path = path
            self.dd_file_name.config(text=os.path.basename(path), fg=self.TEXT)

    def go_style_page(self):
        if not self.jushuitan_path or not self.doudian_path:
            messagebox.showwarning("缺少文件", "请先选择聚水潭和抖店两个 Excel 文件。")
            return
        try:
            js_required = ["款式编码", "颜色", "规格", "商品编码"]
            dd_required = ["商品id(必填)", "商品规格", "商家编码(必填)"]
            self.js_header_idx = self._detect_header_row(self.jushuitan_path, js_required)
            self.dd_header_idx = self._detect_header_row(self.doudian_path, dd_required)
            self.jushuitan_df = pd.read_excel(self.jushuitan_path, header=self.js_header_idx, dtype=object)
            self.doudian_df = pd.read_excel(self.doudian_path, header=self.dd_header_idx, dtype=object)
            self.jushuitan_df.columns = [self.norm_text(c) for c in self.jushuitan_df.columns]
            self.doudian_df.columns = [self.norm_text(c) for c in self.doudian_df.columns]
        except Exception as e:
            messagebox.showerror("读取失败", f"文件读取或表头识别失败：\n{e}")
            return

        for col in ["款式编码", "颜色", "规格", "商品编码"]:
            if col not in self.jushuitan_df.columns:
                messagebox.showerror("格式错误", f"聚水潭文件缺少必填列：{col}")
                return
        for col in ["商品id(必填)", "商品规格", "商家编码(必填)"]:
            if col not in self.doudian_df.columns:
                messagebox.showerror("格式错误", f"抖店文件缺少必填列：{col}")
                return

        # 统一关键字段，避免 pandas/openpyxl 读取后的类型差异。
        self.jushuitan_df["款式编码"] = self.jushuitan_df["款式编码"].map(self.norm_id)
        self.jushuitan_df["颜色"] = self.jushuitan_df["颜色"].map(self.norm_text)
        self.jushuitan_df["规格"] = self.jushuitan_df["规格"].map(self.norm_text)
        self.jushuitan_df["商品编码"] = self.jushuitan_df["商品编码"].map(self.norm_id)
        self.doudian_df["商品id(必填)"] = self.doudian_df["商品id(必填)"].map(self.norm_id)
        self.doudian_df["商品规格"] = self.doudian_df["商品规格"].map(self.norm_text)

        self.style_list = list(dict.fromkeys(v for v in self.jushuitan_df["款式编码"].tolist() if v))
        self.dd_pid_list = list(dict.fromkeys(v for v in self.doudian_df["商品id(必填)"].tolist() if v))
        if not self.style_list:
            messagebox.showerror("数据为空", "聚水潭文件中未找到有效的款式编码。")
            return
        if not self.dd_pid_list:
            messagebox.showerror("数据为空", "抖店文件中未找到有效的商品 ID。")
            return

        self.show_style_page(reset=True)

    # ----------------------------- 第二步 -----------------------------
    def show_style_page(self, reset=False):
        self._clear_page()
        self._page_header(
            "步骤 2 / 3",
            "匹配聚水潭款式与抖店 PID",
            "只需为要处理的聚水潭款式选择 PID；不处理的款式留空即可跳过。同一个抖店 PID 可供多个款式使用，之后按颜色和尺码定位 SKU。"
        )

        if reset or not self.style_pid_rows:
            self.style_pid_rows = {s: [] for s in self.style_list}
        self.style_card_map = {}

        tools = tk.Frame(self.page, bg=self.BG)
        tools.pack(fill="x", pady=(0, 10))
        ttk.Button(tools, text="← 返回文件选择", style="Ghost.TButton", command=self.show_file_page).pack(side="left")
        ttk.Button(tools, text="导入款号对照表", style="Ghost.TButton", command=self.import_mapping_table).pack(side="right")
        self.style_summary_label = tk.Label(tools, text="", bg=self.BG, fg=self.MUTED, font=("微软雅黑", 9))
        self.style_summary_label.pack(side="right", padx=14)

        scroller, _canvas, inner = self._make_scroller(self.page)
        scroller.pack(fill="both", expand=True)

        # 若是从颜色页返回，保留已选择的 style_to_pid。
        previous = {} if reset else (self.style_to_pid.copy() if self.style_to_pid else {})
        for style_code in self.style_list:
            self._build_style_card(inner, style_code)
            pids = previous.get(style_code, [])
            if pids:
                self._set_style_pids(style_code, pids)
            else:
                self._add_pid_row(style_code)

        bottom = tk.Frame(self.page, bg=self.BG)
        bottom.pack(fill="x", pady=(14, 0))
        ttk.Button(bottom, text="下一步：颜色配对", style="Primary.TButton", command=self.go_color_page).pack(side="right")
        self._update_style_summary()

    def _build_style_card(self, parent, style_code):
        card = tk.Frame(parent, bg=self.CARD, highlightbackground=self.BORDER, highlightthickness=1)
        card.pack(fill="x", pady=5, padx=1)
        header = tk.Frame(card, bg=self.CARD)
        header.pack(fill="x", padx=14, pady=(10, 6))
        tk.Label(header, text=style_code, bg=self.CARD, fg=self.TEXT, font=("微软雅黑", 10, "bold")).pack(side="left")
        tk.Label(header, text="聚水潭款式（PID 留空即跳过）", bg=self.CARD, fg=self.MUTED, font=("微软雅黑", 8)).pack(side="left", padx=8)
        ttk.Button(header, text="+ 添加 PID", style="Small.TButton", command=lambda s=style_code: self._add_pid_row(s)).pack(side="right")

        rows_frame = tk.Frame(card, bg=self.CARD)
        rows_frame.pack(fill="x", padx=14, pady=(0, 10))
        self.style_card_map[style_code] = {"card": card, "rows_frame": rows_frame}
        self.style_pid_rows[style_code] = []

    def _add_pid_row(self, style_code, value=""):
        rows_frame = self.style_card_map[style_code]["rows_frame"]
        row_frame = tk.Frame(rows_frame, bg="#FAFAFA")
        row_frame.pack(fill="x", pady=3)

        idx_label = tk.Label(row_frame, text="PID", width=5, bg="#FAFAFA", fg=self.MUTED, font=("微软雅黑", 9))
        idx_label.pack(side="left", padx=(8, 4), pady=5)
        cb = ttk.Combobox(row_frame, values=self.dd_pid_list, font=("微软雅黑", 10), width=42)
        cb.pack(side="left", fill="x", expand=True, padx=4, pady=5)
        if value:
            cb.set(value)
        remove_btn = ttk.Button(row_frame, text="删除", style="Small.TButton")
        remove_btn.pack(side="right", padx=6, pady=4)

        info = {"frame": row_frame, "combobox": cb, "label": idx_label, "remove": remove_btn}
        self.style_pid_rows[style_code].append(info)
        remove_btn.config(command=lambda s=style_code, item=info: self._remove_pid_row(s, item))
        cb.bind("<<ComboboxSelected>>", lambda _e: self._update_style_summary())
        cb.bind("<KeyRelease>", lambda _e: self._update_style_summary())
        self._renumber_pid_rows(style_code)
        self._update_style_summary()
        return info

    def _remove_pid_row(self, style_code, info):
        rows = self.style_pid_rows[style_code]
        if len(rows) <= 1:
            messagebox.showinfo("提示", "每个聚水潭款式至少保留一条 PID。")
            return
        if info in rows:
            rows.remove(info)
            info["frame"].destroy()
            self._renumber_pid_rows(style_code)
            self._update_style_summary()

    def _renumber_pid_rows(self, style_code):
        rows = self.style_pid_rows.get(style_code, [])
        for i, info in enumerate(rows, start=1):
            info["label"].config(text=f"PID {i}")

    def _set_style_pids(self, style_code, pids):
        for info in self.style_pid_rows.get(style_code, []):
            info["frame"].destroy()
        self.style_pid_rows[style_code] = []
        for pid in pids or [""]:
            self._add_pid_row(style_code, pid)

    def _collect_current_style_pid(self, validate=False):
        result = {}
        for style in self.style_list:
            pids = []
            for info in self.style_pid_rows.get(style, []):
                pid = self.norm_id(info["combobox"].get())
                if not pid:
                    # 留空代表此款式/PID 行不参与匹配，允许只处理部分款式。
                    continue
                if pid not in self.dd_pid_list:
                    if validate:
                        raise ValueError(f"款式【{style}】中的 PID【{pid}】不在抖店文件中。")
                    continue
                if pid in pids:
                    if validate:
                        raise ValueError(f"款式【{style}】重复选择了 PID【{pid}】。")
                    continue
                pids.append(pid)
            result[style] = pids
        return result

    def _update_style_summary(self):
        if not hasattr(self, "style_summary_label") or not self.style_summary_label.winfo_exists():
            return
        current = self._collect_current_style_pid(validate=False)
        pid_to_styles = defaultdict(set)
        selected = 0
        for style, pids in current.items():
            selected += len(pids)
            for pid in pids:
                pid_to_styles[pid].add(style)
        shared = sum(1 for styles in pid_to_styles.values() if len(styles) > 1)
        active_styles = sum(bool(pids) for pids in current.values())
        self.style_summary_label.config(
            text=f"已选择 {active_styles}/{len(self.style_list)} 个款式 · {selected} 条关系 · 共享 PID {shared} 个"
        )

    def import_mapping_table(self):
        path = filedialog.askopenfilename(filetypes=[("Excel 文件", "*.xlsx *.xls")])
        if not path:
            return
        try:
            df = pd.read_excel(path, dtype=object)
            df.columns = [self.norm_text(c) for c in df.columns]
        except Exception as e:
            messagebox.showerror("读取失败", f"导入表读取出错：\n{e}")
            return

        if "聚水潭款号" not in df.columns:
            messagebox.showerror("格式错误", "导入表缺少列：聚水潭款号")
            return
        pid_cols = [c for c in df.columns if str(c).startswith("抖店款号")]
        if not pid_cols:
            messagebox.showerror("格式错误", "导入表缺少抖店款号列，例如：抖店款号、抖店款号1、抖店款号2。")
            return

        fill_map = defaultdict(list)
        missing_styles = []
        invalid = []
        for _, row in df.iterrows():
            style = self.norm_id(row.get("聚水潭款号"))
            if not style:
                continue
            if style not in self.style_list:
                if style not in missing_styles:
                    missing_styles.append(style)
                continue
            for col in pid_cols:
                pid = self.norm_id(row.get(col))
                if not pid:
                    continue
                if pid not in fill_map[style]:
                    fill_map[style].append(pid)
                if pid not in self.dd_pid_list:
                    invalid.append(f"{style} → {pid}")

        if not fill_map:
            messagebox.showwarning("没有匹配", "导入表中没有找到可用于当前聚水潭文件的款式关系。")
            return

        # 只替换导入表中出现的款式，未出现的款式保留当前选择。
        for style, pids in fill_map.items():
            self._set_style_pids(style, pids)
        self._update_style_summary()

        lines = [f"已填充 {len(fill_map)} 个聚水潭款式。", "同一个抖店款号重复出现在不同聚水潭款式中时会保留，不再覆盖。"]
        if missing_styles:
            lines.append("\n聚水潭文件中不存在，已跳过：\n" + "、".join(missing_styles[:20]))
        if invalid:
            lines.append("\n抖店文件中不存在的 PID（已填入，请修正）：\n" + "、".join(invalid[:20]))
        messagebox.showinfo("导入完成", "\n".join(lines))

    def go_color_page(self):
        try:
            self.style_to_pid = self._collect_current_style_pid(validate=True)
        except ValueError as e:
            messagebox.showwarning("请检查 PID", str(e))
            return
        if not any(self.style_to_pid.values()):
            messagebox.showwarning("尚未选择款式", "请至少为一个需要匹配的聚水潭款式选择抖店 PID；其他款式可以留空。")
            return
        self.show_color_page()

    # ----------------------------- 第三步 -----------------------------
    def show_color_page(self):
        self._clear_page()
        self._page_header(
            "步骤 3 / 3",
            "颜色配对与生成",
            "颜色相同会自动配对；其他颜色按需手动配对即可。无需配齐全部颜色：只有已配对颜色对应的尺码会写入编码，未配对的颜色会跳过并保留原商家编码。共享 PID 的同一抖店颜色仍只能匹配一次。"
        )

        top_actions = tk.Frame(self.page, bg=self.BG)
        top_actions.pack(fill="x", pady=(0, 10))
        ttk.Button(top_actions, text="← 返回 PID 配对", style="Ghost.TButton", command=self._back_to_style_page).pack(side="left")
        self.color_summary_label = tk.Label(top_actions, text="", bg=self.BG, fg=self.MUTED, font=("微软雅黑", 9))
        self.color_summary_label.pack(side="right", padx=4)

        scroller, _canvas, inner = self._make_scroller(self.page)
        scroller.pack(fill="both", expand=True)

        self.color_mapping = {}
        self.temp_selected = {}
        self.js_listbox_map = {}
        self.dd_listbox_map = {}
        self.pair_label_map = {}
        self.block_meta = {}

        # 先建立全部映射容器，再做自动配对，才能正确识别共享 PID 的占用关系。
        for style in self.style_list:
            self.color_mapping[style] = {}
            for pid in self.style_to_pid[style]:
                self.color_mapping[style][pid] = {}
                self.temp_selected[(style, pid)] = None

        for style in self.style_list:
            for pid in self.style_to_pid[style]:
                self._build_color_block(inner, style, pid)

        # 自动配对：相同色名且该 PID 的颜色尚未被其他款式占用。
        for style in self.style_list:
            for pid in self.style_to_pid[style]:
                meta = self.block_meta[(style, pid)]
                for js_color in meta["js_colors"]:
                    if js_color in meta["dd_colors"] and not self._find_dd_color_owner(pid, js_color):
                        self.color_mapping[style][pid][js_color] = js_color

        for pid in set(p for ps in self.style_to_pid.values() for p in ps):
            self.refresh_pid_visuals(pid)

        bottom = tk.Frame(self.page, bg=self.BG)
        bottom.pack(fill="x", pady=(14, 0))
        ttk.Button(bottom, text="检查并生成匹配文件", style="Primary.TButton", command=self.process_and_save).pack(side="right")

    def _back_to_style_page(self):
        # style_to_pid 已保存，show_style_page 会恢复选择。
        self.show_style_page(reset=False)

    def _build_color_block(self, parent, style_code, pid):
        js_colors = list(dict.fromkeys(
            c for c in self.jushuitan_df.loc[self.jushuitan_df["款式编码"] == style_code, "颜色"].tolist() if c
        ))
        dd_colors = self.get_dd_color_list(pid)
        key = (style_code, pid)
        self.block_meta[key] = {"js_colors": js_colors, "dd_colors": dd_colors}

        frame = ttk.LabelFrame(parent, text=f"聚水潭款式：{style_code}    ·    抖店 PID：{pid}", padding=10)
        frame.pack(fill="x", padx=1, pady=7)

        shared_styles = [s for s, pids in self.style_to_pid.items() if pid in pids]
        if len(shared_styles) > 1:
            tk.Label(
                frame,
                text=f"共享 PID：此 PID 同时关联 {len(shared_styles)} 个聚水潭款式（{'、'.join(shared_styles)}）",
                bg="#FFF7E6", fg="#B88230", font=("微软雅黑", 8), anchor="w"
            ).pack(fill="x", pady=(0, 8), ipady=4, padx=1)

        grid = tk.Frame(frame, bg=self.CARD)
        grid.pack(fill="x")
        for col in range(3):
            grid.columnconfigure(col, weight=1, uniform="color_cols")

        left = tk.Frame(grid, bg=self.CARD)
        left.grid(row=0, column=0, sticky="nsew", padx=(0, 6))
        tk.Label(left, text="聚水潭颜色", bg="#F0F9EB", fg="#529B2E", font=("微软雅黑", 9, "bold"), pady=6).pack(fill="x")
        js_lb = tk.Listbox(left, height=max(6, min(10, len(js_colors) + 1)), exportselection=False,
                           font=("微软雅黑", 9), bd=1, relief="solid", highlightthickness=0,
                           selectbackground=self.PRIMARY, selectforeground="#FFFFFF")
        for color in js_colors:
            js_lb.insert(tk.END, color)
        js_lb.pack(fill="both", expand=True, pady=(5, 0))
        self.js_listbox_map[key] = js_lb

        mid = tk.Frame(grid, bg=self.CARD)
        mid.grid(row=0, column=1, sticky="nsew", padx=6)
        tk.Label(mid, text="当前映射", bg="#F4F4F5", fg="#606266", font=("微软雅黑", 9, "bold"), pady=6).pack(fill="x")
        pair_label = tk.Label(mid, text="", bg=self.CARD, fg=self.TEXT, justify="left", anchor="nw", font=("微软雅黑", 9), wraplength=300)
        pair_label.pack(fill="both", expand=True, pady=7)
        self.pair_label_map[key] = pair_label
        actions = tk.Frame(mid, bg=self.CARD)
        actions.pack(fill="x", pady=(4, 0))
        ttk.Button(actions, text="撤销选中", style="Small.TButton", command=lambda s=style_code, p=pid: self.unpair_selected(s, p)).pack(side="left")
        ttk.Button(actions, text="重置此块", style="Small.TButton", command=lambda s=style_code, p=pid: self.reset_one_style_pid(s, p)).pack(side="right")

        right = tk.Frame(grid, bg=self.CARD)
        right.grid(row=0, column=2, sticky="nsew", padx=(6, 0))
        tk.Label(right, text="抖店颜色", bg="#ECF5FF", fg="#337ECC", font=("微软雅黑", 9, "bold"), pady=6).pack(fill="x")
        dd_lb = tk.Listbox(right, height=max(6, min(10, len(dd_colors) + 1)), exportselection=False,
                           font=("微软雅黑", 9), bd=1, relief="solid", highlightthickness=0,
                           selectbackground=self.PRIMARY, selectforeground="#FFFFFF")
        for color in dd_colors:
            dd_lb.insert(tk.END, color)
        dd_lb.pack(fill="both", expand=True, pady=(5, 0))
        self.dd_listbox_map[key] = dd_lb

        js_lb.bind("<<ListboxSelect>>", lambda _e, s=style_code, p=pid: self.on_js_click(s, p))
        dd_lb.bind("<<ListboxSelect>>", lambda _e, s=style_code, p=pid: self.on_dd_click(s, p))

    def get_dd_color_list(self, product_id):
        rows = self.doudian_df[self.doudian_df["商品id(必填)"] == self.norm_id(product_id)]
        colors = []
        for spec in rows["商品规格"].tolist():
            color = self.extract_color(spec)
            if color and color not in colors:
                colors.append(color)
        return colors

    @staticmethod
    def _extract_spec_value(spec_str, accepted_keys):
        if spec_str is None:
            return None
        text = str(spec_str).strip().strip("{}")
        parts = re.split(r"[，,；;]", text)
        accepted = {k.strip() for k in accepted_keys}
        for part in parts:
            part = part.strip().strip("'\"")
            if not part:
                continue
            kv = re.split(r"[:：]", part, maxsplit=1)
            if len(kv) != 2:
                continue
            key = kv[0].strip().strip("'\"")
            value = kv[1].strip().strip("'\"")
            if key in accepted:
                return value or None
        return None

    def extract_color(self, spec_str):
        return self._extract_spec_value(spec_str, {"颜色分类", "颜色", "Color", "color"})

    def extract_size(self, spec_str):
        return self._extract_spec_value(spec_str, {"尺码大小", "尺码", "规格", "Size", "size"})

    def _find_dd_color_owner(self, pid, dd_color, exclude=None):
        """返回同一 PID 下某抖店颜色的占用者 (style, js_color)，没有则 None。"""
        for style, pid_maps in self.color_mapping.items():
            if pid not in pid_maps:
                continue
            for js_color, mapped_dd in pid_maps[pid].items():
                if mapped_dd == dd_color:
                    owner = (style, js_color)
                    if exclude and owner == exclude:
                        continue
                    return owner
        return None

    def on_js_click(self, style_code, pid):
        key = (style_code, pid)
        lb = self.js_listbox_map[key]
        selection = lb.curselection()
        if not selection:
            return
        self.temp_selected[key] = lb.get(selection[0])

    def on_dd_click(self, style_code, pid):
        key = (style_code, pid)
        js_color = self.temp_selected.get(key)
        dd_lb = self.dd_listbox_map[key]
        selection = dd_lb.curselection()
        if not selection:
            return
        dd_color = dd_lb.get(selection[0])

        if not js_color:
            dd_lb.selection_clear(0, tk.END)
            messagebox.showinfo("操作提示", "请先点击左侧聚水潭颜色，再点击右侧抖店颜色。")
            return

        current_map = self.color_mapping[style_code][pid]
        if js_color in current_map:
            messagebox.showinfo("已配对", f"聚水潭颜色【{js_color}】已经配对。需要修改时请先“撤销选中”。")
        else:
            owner = self._find_dd_color_owner(pid, dd_color)
            if owner:
                owner_style, owner_js = owner
                messagebox.showwarning(
                    "颜色已被占用",
                    f"抖店颜色【{dd_color}】在 PID【{pid}】下已经分配给：\n"
                    f"聚水潭款式【{owner_style}】 / 颜色【{owner_js}】。\n\n"
                    "共享 PID 场景下，同一个抖店颜色只能归属一个聚水潭款式。"
                )
            else:
                current_map[js_color] = dd_color
                self.refresh_pid_visuals(pid)

        self.temp_selected[key] = None
        dd_lb.selection_clear(0, tk.END)
        self.js_listbox_map[key].selection_clear(0, tk.END)

    def unpair_selected(self, style_code, pid):
        key = (style_code, pid)
        lb = self.js_listbox_map[key]
        selection = lb.curselection()
        if not selection:
            messagebox.showinfo("提示", "请先在左侧选中一个已经配对的聚水潭颜色。")
            return
        js_color = lb.get(selection[0])
        mapping = self.color_mapping[style_code][pid]
        if js_color not in mapping:
            messagebox.showinfo("提示", f"颜色【{js_color}】当前没有配对。")
            return
        del mapping[js_color]
        self.temp_selected[key] = None
        self.refresh_pid_visuals(pid)

    def reset_one_style_pid(self, style_code, pid):
        self.color_mapping[style_code][pid] = {}
        self.temp_selected[(style_code, pid)] = None
        self.refresh_pid_visuals(pid)

    def refresh_pid_visuals(self, pid):
        """共享 PID 的任意映射变化后，刷新所有相关块的颜色占用状态。"""
        for style in self.style_list:
            key = (style, pid)
            if key not in self.block_meta:
                continue
            mapping = self.color_mapping[style][pid]
            js_lb = self.js_listbox_map[key]
            dd_lb = self.dd_listbox_map[key]

            for i in range(js_lb.size()):
                js_lb.itemconfig(i, bg="#FFFFFF", fg=self.TEXT)
            for js_color in mapping:
                try:
                    idx = list(js_lb.get(0, tk.END)).index(js_color)
                    js_lb.itemconfig(idx, bg="#F0F9EB", fg="#529B2E")
                except ValueError:
                    pass

            for i in range(dd_lb.size()):
                dd_color = dd_lb.get(i)
                owner = self._find_dd_color_owner(pid, dd_color)
                if not owner:
                    dd_lb.itemconfig(i, bg="#FFFFFF", fg=self.TEXT)
                elif owner[0] == style:
                    dd_lb.itemconfig(i, bg="#ECF5FF", fg="#337ECC")
                else:
                    dd_lb.itemconfig(i, bg="#F4F4F5", fg=self.MUTED)

            text_lines = [f"{js}  →  {dd}" for js, dd in mapping.items()]
            total = len(self.block_meta[key]["js_colors"])
            header = f"已配对 {len(mapping)} / {total}（其余跳过）"
            self.pair_label_map[key].config(text=header + ("\n\n" + "\n".join(text_lines) if text_lines else "\n\n未配对"))
        self._update_color_summary()

    def _update_color_summary(self):
        if not hasattr(self, "color_summary_label") or not self.color_summary_label.winfo_exists():
            return
        total = 0
        paired = 0
        for style in self.style_list:
            js_colors = list(dict.fromkeys(
                c for c in self.jushuitan_df.loc[self.jushuitan_df["款式编码"] == style, "颜色"].tolist() if c
            ))
            for pid in self.style_to_pid[style]:
                total += len(js_colors)
                paired += len(self.color_mapping.get(style, {}).get(pid, {}))
        self.color_summary_label.config(text=f"颜色配对 {paired} / {total} · 未配对颜色不导出修改")

    # ----------------------------- 导出 -----------------------------
    def _validate_color_mapping(self):
        """只要求至少存在一组有效配对；未配对款式与颜色均允许跳过。"""
        paired_count = 0
        conflicts = []
        for pid in set(p for pids in self.style_to_pid.values() for p in pids):
            owners = defaultdict(list)
            for style in self.style_list:
                if pid not in self.color_mapping.get(style, {}):
                    continue
                mapping = self.color_mapping[style][pid]
                js_colors = set(self.block_meta.get((style, pid), {}).get("js_colors", []))
                dd_colors = set(self.block_meta.get((style, pid), {}).get("dd_colors", []))
                for js_color, dd_color in mapping.items():
                    # 二次校验手工配对，防止异常映射写错 SKU。
                    if js_color not in js_colors or dd_color not in dd_colors:
                        raise ValueError(
                            f"无效颜色配对：款式【{style}】/ PID【{pid}】/ "
                            f"聚水潭【{js_color}】→ 抖店【{dd_color}】。请重新配对。"
                        )
                    paired_count += 1
                    owners[dd_color].append((style, js_color))
            for dd_color, items in owners.items():
                if len(items) > 1:
                    conflicts.append(f"PID {pid} / 抖店颜色 {dd_color}: {items}")

        if conflicts:
            preview = "\n".join(conflicts[:8])
            raise ValueError("检测到共享 PID 颜色重复占用，请先修正：\n\n" + preview)
        if paired_count == 0:
            raise ValueError("尚未配对任何颜色。请至少配对一组需要导出的颜色；其他颜色无需配对。")

    def _build_js_lookup(self):
        """(style, color, size) -> [商品编码]，保留重复键以便检查歧义。"""
        lookup = defaultdict(list)
        for _, row in self.jushuitan_df.iterrows():
            style = self.norm_id(row.get("款式编码"))
            color = self.norm_text(row.get("颜色"))
            size = self.norm_text(row.get("规格"))
            code = self.norm_id(row.get("商品编码"))
            if not (style and color and size and code):
                continue
            key = (style, color, size)
            if code not in lookup[key]:
                lookup[key].append(code)
        return lookup

    def _resolve_code(self, pid, dd_color, dd_size, js_lookup):
        """共享 PID 安全解析。返回 (status, code, detail)。"""
        candidate_codes = []
        details = []
        has_color_pair = False
        candidate_styles = [s for s, pids in self.style_to_pid.items() if pid in pids]

        for style in candidate_styles:
            mapping = self.color_mapping.get(style, {}).get(pid, {})
            reverse = {dd: js for js, dd in mapping.items()}
            if dd_color not in reverse:
                continue
            has_color_pair = True
            js_color = reverse[dd_color]
            codes = js_lookup.get((style, js_color, dd_size), [])
            if not codes:
                details.append(f"{style}/{js_color}/{dd_size} 无聚水潭商品编码")
                continue
            for code in codes:
                if code not in candidate_codes:
                    candidate_codes.append(code)
                    details.append(f"{style}/{js_color}/{dd_size} → {code}")

        if len(candidate_codes) == 1:
            return "ok", candidate_codes[0], details
        if len(candidate_codes) > 1:
            return "ambiguous", None, details
        if not has_color_pair:
            return "unpaired", None, details
        return "missing", None, details

    def _load_workbook_for_export(self):
        suffix = Path(self.doudian_path).suffix.lower()
        if suffix in {".xlsx", ".xlsm"}:
            wb = load_workbook(self.doudian_path, keep_vba=(suffix == ".xlsm"))
            return wb, wb.active, False

        # .xls 无法由 openpyxl 原样加载：转成新的 xlsx 工作簿，保留单元格数据。
        raw = pd.read_excel(self.doudian_path, header=None, dtype=object)
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        for r_idx, row in enumerate(raw.itertuples(index=False, name=None), start=1):
            for c_idx, value in enumerate(row, start=1):
                if pd.isna(value):
                    value = None
                ws.cell(r_idx, c_idx, value=value)
        return wb, ws, True

    def _locate_export_columns(self, ws):
        header_row = self.dd_header_idx + 1
        header_map = {}
        for col in range(1, ws.max_column + 1):
            name = self.norm_text(ws.cell(header_row, col).value)
            if name:
                header_map[name] = col
        needed = ["商品id(必填)", "商品规格", "商家编码(必填)"]
        missing = [x for x in needed if x not in header_map]
        if missing:
            raise ValueError("抖店文件导出时找不到列：" + "、".join(missing))
        return header_row, header_map

    def process_and_save(self):
        try:
            self._validate_color_mapping()
        except ValueError as e:
            messagebox.showwarning("配对未完成", str(e))
            return

        try:
            js_lookup = self._build_js_lookup()
            wb, ws, converted_from_xls = self._load_workbook_for_export()
            header_row, cols = self._locate_export_columns(ws)

            mapped_pids = set(p for pids in self.style_to_pid.values() for p in pids)
            stats = defaultdict(int)
            problem_rows = []

            for row_idx in range(header_row + 1, ws.max_row + 1):
                pid = self.norm_id(ws.cell(row_idx, cols["商品id(必填)"]).value)
                if not pid or pid not in mapped_pids:
                    continue
                stats["target_rows"] += 1
                spec = ws.cell(row_idx, cols["商品规格"]).value
                dd_color = self.extract_color(spec)
                if not dd_color:
                    stats["parse_failed"] += 1
                    problem_rows.append(f"第 {row_idx} 行：无法解析颜色 | PID {pid} | {spec}")
                    continue

                # 仅处理已经配对的抖店颜色，未配对颜色无需检查尺码或补齐。
                if not any(dd_color in mapping.values()
                           for mapping in (self.color_mapping.get(style, {}).get(pid, {})
                                           for style in self.style_list)):
                    stats["skipped_unpaired"] += 1
                    continue

                dd_size = self.extract_size(spec)
                if not dd_size:
                    stats["parse_failed"] += 1
                    problem_rows.append(f"第 {row_idx} 行：已配对颜色缺少尺码 | PID {pid} | {spec}")
                    continue

                status, code, detail = self._resolve_code(pid, dd_color, dd_size, js_lookup)
                if status == "ok":
                    ws.cell(row_idx, cols["商家编码(必填)"], value=code)
                    stats["matched"] += 1
                elif status == "ambiguous":
                    stats["ambiguous"] += 1
                    problem_rows.append(f"第 {row_idx} 行：存在多个候选编码 | PID {pid} | {dd_color}/{dd_size} | {'; '.join(detail)}")
                elif status == "unpaired":
                    stats["skipped_unpaired"] += 1
                else:
                    stats["missing"] += 1
                    problem_rows.append(f"第 {row_idx} 行：未找到聚水潭 SKU | PID {pid} | {dd_color}/{dd_size} | {'; '.join(detail)}")

            # 多候选属于明确的数据歧义，为了“正确优先”不允许静默导出错误结果。
            if stats["ambiguous"]:
                preview = "\n".join(problem_rows[:10])
                messagebox.showerror(
                    "发现歧义，已停止生成",
                    f"有 {stats['ambiguous']} 行同时匹配到多个聚水潭编码。为避免写错，程序没有保存文件。\n\n{preview}"
                )
                return

            summary = (
                f"关联 PID 的 SKU 行：{stats['target_rows']}\n"
                f"成功写入：{stats['matched']}\n"
                f"未配对颜色（已跳过、原值不变）：{stats['skipped_unpaired']}\n"
                f"已配对但未找到 SKU：{stats['missing']}\n"
                f"规格解析失败：{stats['parse_failed']}"
            )
            if stats["missing"] or stats["parse_failed"]:
                preview = "\n".join(problem_rows[:8])
                go_on = messagebox.askyesno(
                    "存在未匹配行",
                    summary + "\n\n有已配对颜色的行无法写入编码，这些行将保留原商家编码。是否仍然生成文件？\n\n" + preview
                )
                if not go_on:
                    return

            if stats["matched"] == 0:
                messagebox.showwarning("没有成功匹配", summary + "\n\n没有任何 SKU 成功匹配，未生成文件。请检查配对颜色和尺码。")
                return

            save_path = filedialog.asksaveasfilename(
                defaultextension=".xlsx",
                filetypes=[("Excel 文件", "*.xlsx")],
                initialfile="抖店编码_已匹配.xlsx"
            )
            if not save_path:
                return
            wb.save(save_path)

            extra = "\n注意：原抖店文件为 .xls，已转换为 .xlsx，单元格样式可能与原文件略有差异。" if converted_from_xls else ""
            messagebox.showinfo("处理完成", f"文件已保存：\n{save_path}\n\n{summary}{extra}")
        except Exception as e:
            messagebox.showerror("处理失败", f"生成文件时出现错误：\n{e}")
            traceback.print_exc()


if __name__ == "__main__":
    root = tk.Tk()
    app = ExcelMatcherApp(root)
    root.mainloop()
