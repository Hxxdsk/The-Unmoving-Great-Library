from __future__ import annotations

import csv
import hashlib
import json
import re
import shutil
import unicodedata
import uuid
from collections import OrderedDict
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Callable

from openpyxl import load_workbook

IMAGE_SUFFIXES = {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".tif", ".tiff", ".gif"}
REQUIRED_COLUMNS = ("供应商款式编码", "颜色", "商品编码")


class ValidationError(ValueError):
    pass


@dataclass
class ImageMatch:
    filename: str
    color_label: str
    main_color: str
    confidence: float
    needs_review: bool
    reason: str = ""


@dataclass
class Style:
    code: str
    colors: dict[str, list[str]] = field(default_factory=OrderedDict)
    folder: Path | None = None
    chart: Path | None = None
    images: list[Path] = field(default_factory=list)
    issues: list[str] = field(default_factory=list)
    chart_colors: list[str] = field(default_factory=list)
    matches: dict[str, ImageMatch] = field(default_factory=dict)
    selections: dict[str, str] = field(default_factory=dict)
    manual_colors: set[str] = field(default_factory=set)
    analysis_error: str = ""

    @property
    def sku_count(self) -> int:
        return sum(len(codes) for codes in self.colors.values())


@dataclass
class Project:
    root: Path
    workbook: Path
    styles: list[Style]
    warnings: list[str] = field(default_factory=list)

    @property
    def sku_count(self) -> int:
        return sum(style.sku_count for style in self.styles)


def text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalized(value: str) -> str:
    return unicodedata.normalize("NFKC", value).strip().casefold()


def color_key(value: str) -> str:
    value = re.sub(r"\s+", "", normalized(value))
    return value[:-1] if value.endswith("色") else value


def safe_component(value: str) -> str:
    if (not value or value in {".", ".."} or len(value) > 180
            or re.search(r'[<>:"/\\|?*\x00-\x1f]', value)
            or value.endswith((".", " "))):
        raise ValidationError(f"不能用作 Windows 文件名的编码：{value!r}")
    base = value.split(".")[0].upper()
    if re.fullmatch(r"CON|PRN|AUX|NUL|COM[1-9]|LPT[1-9]", base):
        raise ValidationError(f"编码属于 Windows 保留文件名：{value}")
    return value


def natural_key(path: Path):
    return [int(p) if p.isdigit() else p.casefold() for p in re.split(r"(\d+)", path.name)]


def list_workbooks(root: Path) -> list[Path]:
    return sorted((p for p in root.iterdir() if p.is_file()
                   and p.suffix.lower() in {".xlsx", ".xlsm"}
                   and not p.name.startswith("~$")), key=lambda p: p.name)


def read_project(root: Path, workbook: Path) -> Project:
    root, workbook = root.resolve(), workbook.resolve()
    if workbook.parent != root:
        raise ValidationError("请使用所选目录内的表格。")
    wb = load_workbook(workbook, read_only=True, data_only=True)
    styles: dict[str, Style] = OrderedDict()
    seen_skus: dict[str, tuple[str, str]] = {}
    warnings = []
    found = False
    try:
        for sheet in wb.worksheets:
            # 部分 ERP 导出器把实际 25 列 / 数十行的尺寸错误写为 A1:A1。
            # 只读模式必须忽略该提示，按实际 XML 单元格流读取。
            sheet.reset_dimensions()
            header = None
            for row_no, cells in enumerate(sheet.iter_rows(), 1):
                values = []
                for cell in cells:
                    value = text(cell.value)
                    if isinstance(cell.value, (int, float)) and re.fullmatch(r"0+", cell.number_format):
                        value = value.zfill(len(cell.number_format))
                    values.append(value)
                if header is None:
                    indexes = {v: i for i, v in enumerate(values) if v in REQUIRED_COLUMNS}
                    if set(indexes) == set(REQUIRED_COLUMNS):
                        if any(values.count(name) != 1 for name in REQUIRED_COLUMNS):
                            raise ValidationError(f"{sheet.title} 第 {row_no} 行含重复列名。")
                        header = indexes
                        found = True
                    elif row_no >= 30:
                        break
                    continue
                if not any(values):
                    continue
                code, color, sku = (values[header[name]] if header[name] < len(values) else ""
                                    for name in REQUIRED_COLUMNS)
                if not all((code, color, sku)):
                    raise ValidationError(f"{sheet.title} 第 {row_no} 行的供应商款式编码、颜色或商品编码为空。")
                safe_component(code)
                safe_component(sku)
                identity = (code, color)
                if normalized(sku) in seen_skus:
                    if seen_skus[normalized(sku)] != identity:
                        raise ValidationError(f"商品编码 {sku} 被用于不同款号或颜色，无法安全输出。")
                    warnings.append(f"已合并重复商品编码：{sku}（{sheet.title} 第 {row_no} 行）")
                    continue
                seen_skus[normalized(sku)] = identity
                style = styles.setdefault(code, Style(code))
                style.colors.setdefault(color, []).append(sku)
    finally:
        wb.close()
    if not found:
        raise ValidationError("表格中找不到列：供应商款式编码、颜色、商品编码（检查前 30 行表头）。")
    if not styles:
        raise ValidationError("表格中没有可处理的商品行。")
    folders: dict[str, list[Path]] = {}
    for folder in root.iterdir():
        if folder.is_dir():
            folders.setdefault(normalized(folder.name), []).append(folder)
    for style in styles.values():
        candidates = folders.get(normalized(style.code), [])
        if len(candidates) != 1:
            style.issues.append("找不到同名款号文件夹。" if not candidates else "存在多个同名款号文件夹。")
            continue
        style.folder = candidates[0].resolve()
        if not style.folder.is_relative_to(root):
            style.issues.append("款号文件夹链接到了所选目录以外。")
            continue
        files = sorted((p for p in style.folder.iterdir() if p.is_file()
                        and p.suffix.lower() in IMAGE_SUFFIXES), key=natural_key)
        if any(not p.resolve().is_relative_to(style.folder) for p in files):
            style.issues.append("图片链接到了款号文件夹以外。")
            continue
        charts = [p for p in files if normalized(p.stem) == "颜色"]
        if len(charts) != 1:
            style.issues.append("需要一张名为“颜色”的图片。" if not charts else "有多张“颜色”图片，请仅保留一张。")
        else:
            style.chart = charts[0]
        style.images = [p for p in files if normalized(p.stem) != "颜色"]
        if not style.images:
            style.issues.append("没有待配对的商品图片。")
    return Project(root, workbook, list(styles.values()), warnings)


def apply_suggestions(style: Style) -> None:
    """仅名称明确相同且视觉配对可信时自动连接，保留人工选择。"""
    for color in style.colors:
        if color in style.manual_colors:
            continue
        style.selections.pop(color, None)
        labels = [label for label in style.chart_colors if color_key(label) == color_key(color)]
        if len(labels) != 1:
            continue
        matches = [m for m in style.matches.values()
                   if m.color_label == labels[0] and m.confidence >= 0.88 and not m.needs_review]
        if matches:
            best = max(matches, key=lambda m: m.confidence)
            style.selections[color] = best.filename


@dataclass(frozen=True)
class CopyItem:
    style: str
    color: str
    sku: str
    source: Path
    target: Path


def make_plan(project: Project) -> list[CopyItem]:
    plan = []
    names = set()
    for style in project.styles:
        if style.issues:
            raise ValidationError(f"{style.code}：{'；'.join(style.issues)}")
        allowed = {p.name: p for p in style.images}
        for color, skus in style.colors.items():
            chosen = style.selections.get(color)
            if chosen not in allowed:
                raise ValidationError(f"{style.code} / {color} 还没有选择图片。")
            source = allowed[chosen]
            if not source.is_file() or not source.resolve().is_relative_to(style.folder):
                raise ValidationError(f"原图已移走或链接到目录外：{source.name}")
            for sku in skus:
                target = Path(safe_component(style.code)) / (safe_component(sku) + source.suffix.lower())
                if normalized(str(target)) in names:
                    raise ValidationError(f"存在重复输出文件名：{target}")
                names.add(normalized(str(target)))
                plan.append(CopyItem(style.code, color, sku, source, target))
    return plan


def export_project(project: Project, plan: list[CopyItem], cancel=None,
                   progress: Callable[[int, int], None] = lambda *_: None) -> Path:
    if not plan:
        raise ValidationError("没有可输出的商品图片。")
    # 每次都创建新批次；先在临时目录完成全部复制，再一次性公布结果。
    output = project.root / "_商品编码图片"
    if output.is_symlink():
        raise ValidationError("输出目录不能是链接。")
    output.mkdir(exist_ok=True)
    token = uuid.uuid4().hex[:10]
    batch = output / (datetime.now().strftime("%Y%m%d_%H%M%S") + "_" + token)
    stage = output / (".处理中_" + token)
    stage.mkdir()
    records = []
    try:
        for index, item in enumerate(plan, 1):
            if cancel is not None and cancel.is_set():
                raise ValidationError("已取消输出，临时文件已清理。")
            if not item.source.resolve().is_relative_to(project.root):
                raise ValidationError("原图已链接到所选目录外，请重新扫描。")
            destination = stage / item.target
            destination.parent.mkdir(parents=True, exist_ok=True)
            before = item.source.stat()
            shutil.copy2(item.source, destination)
            after = item.source.stat()
            if (before.st_size, before.st_mtime_ns) != (after.st_size, after.st_mtime_ns):
                raise ValidationError(f"复制期间原图发生变化：{item.source.name}，请重新操作。")
            digest = hashlib.sha256(destination.read_bytes()).hexdigest()
            records.append({"款号": item.style, "表格颜色": item.color, "商品编码": item.sku,
                            "原图": str(item.source), "输出文件": str(item.target), "SHA256": digest})
            progress(index, len(plan))
        if cancel is not None and cancel.is_set():
            raise ValidationError("已取消输出，临时文件已清理。")
        with (stage / "配对记录.csv").open("w", encoding="utf-8-sig", newline="") as f:
            writer = csv.DictWriter(f, fieldnames=["款号", "表格颜色", "商品编码", "原图", "输出文件", "SHA256"])
            writer.writeheader()
            for record in records:
                writer.writerow({k: "'" + v if v.startswith(("=", "+", "-", "@")) else v for k, v in record.items()})
        (stage / "配对记录.json").write_text(json.dumps({"表格": str(project.workbook), "文件": records},
                                                           ensure_ascii=False, indent=2), encoding="utf-8")
        stage.rename(batch)
    except BaseException:
        shutil.rmtree(stage, ignore_errors=True)
        raise
    return batch
