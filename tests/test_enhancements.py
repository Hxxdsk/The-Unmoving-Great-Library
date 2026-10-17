import gc
import io
import json
import threading
import time
import zipfile
from pathlib import Path
from unittest.mock import patch

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from PIL import Image, ImageDraw

from color_tool.data import Style, make_plan, read_project, export_project
from color_tool.gemini import GeminiClient, GeminiError
from color_tool.pairing import one_to_one_suggestions
from suite.cache import ResultCache
from suite.legacy import load_original, common_ttk_theme
from suite.materials import validate_material, classify_materials
from suite.submission import create_input_template, classify_image
from suite.workbook_images import read_images
from suite.standards import validate_standard
from suite.joiner_adapter import assign_sides, recognize_sides, filename_hint


@pytest.fixture(autouse=True)
def isolated_ai_cache(monkeypatch, tmp_path):
    monkeypatch.setattr(
        "suite.cache.cache_directory", lambda: tmp_path / "runtime-cache"
    )


def ai():
    return GeminiClient("fixture-key", "models/gemini-test")


def photo(path, color):
    Image.new("RGB", (64, 64), color).save(path)
    return path


def match(filename, label, confidence=0.98):
    return {
        "filename": filename,
        "color_label": label,
        "main_color": label,
        "confidence": confidence,
        "needs_review": False,
        "reason": "主体颜色相符",
    }


def test_manual_export_without_chart_or_api(tmp_path):
    book = Workbook()
    book.active.append(["供应商款式编码", "颜色", "商品编码"])
    for number in range(6):
        book.active.append(["A", "黑色", f"SKU{number}"])
    file = tmp_path / "资料.xlsx"
    book.save(file)
    folder = tmp_path / "A"
    folder.mkdir()
    source = photo(folder / "2.jpg", "black")
    project = read_project(tmp_path, file)
    style = project.styles[0]
    assert style.ai_issues and not style.issues
    style.selections["黑色"] = "2.jpg"
    style.manual_colors.add("黑色")
    result = export_project(project, make_plan(project))
    outputs = list((result / "A").glob("*.jpg"))
    assert len(outputs) == 6 and all(
        output.read_bytes() == source.read_bytes() for output in outputs
    )


def test_partial_manual_export_requires_explicit_option(tmp_path):
    book = Workbook()
    book.active.append(["供应商款式编码", "颜色", "商品编码"])
    book.active.append(["A", "黑色", "BLACK"])
    book.active.append(["A", "白色", "WHITE"])
    path = tmp_path / "table.xlsx"
    book.save(path)
    folder = tmp_path / "A"
    folder.mkdir()
    photo(folder / "1.jpg", "black")
    project = read_project(tmp_path, path)
    project.styles[0].selections["黑色"] = "1.jpg"
    with pytest.raises(ValueError):
        make_plan(project)
    plan = make_plan(project, partial=True)
    assert len(plan) == 1 and plan[0].sku == "BLACK"


def test_exact_pair_locked_then_alias_prefilled_and_cached(tmp_path):
    chart = photo(tmp_path / "颜色.png", "white")
    black, gray = photo(tmp_path / "1.jpg", "black"), photo(tmp_path / "2.jpg", "gray")
    style = Style(
        "A",
        colors={"黑色": ["secret-sku"], "深灰色": ["sku2"]},
        folder=tmp_path,
        chart=chart,
        images=[black, gray],
    )
    reply = {
        "colors": [
            {"label": "黑色", "description": "黑"},
            {"label": "花灰色", "description": "深灰"},
        ],
        "images": [match("1.jpg", "黑色"), match("2.jpg", "花灰色")],
        "alternatives": [
            {
                "table_color": "深灰色",
                "filename": "1.jpg",
                "chart_label": "黑色",
                "confidence": 0.99,
                "reason": "不可挪用",
            },
            {
                "table_color": "深灰色",
                "filename": "2.jpg",
                "chart_label": "花灰色",
                "confidence": 0.72,
                "reason": "主体是深灰",
            },
        ],
    }
    client = ai()
    cache = ResultCache("test", tmp_path / "cache")
    with patch.object(client, "generate", return_value=reply) as request:
        client.analyze_style(style, cache=cache)
        client.analyze_style(style, cache=cache)
        assert request.call_count == 1
        assert "secret-sku" not in json.dumps(request.call_args.args[0])
        client.analyze_style(style, cache=cache, force=True)
        assert request.call_count == 2
    assert style.selections == {"黑色": "1.jpg", "深灰色": "2.jpg"}
    assert style.name_mismatches == {"深灰色": "花灰色"}
    assert style.pair_confidences["深灰色"] == 0.72
    assert "名称不一致" in style.review_notes["深灰色"]
    result = export_project(
        type("P", (), {"root": tmp_path, "workbook": tmp_path / "table.xlsx"})(),
        make_plan(type("P", (), {"styles": [style]})()),
    )
    record = json.loads((result / "配对记录.json").read_text())["文件"][1]
    assert "名称不一致" in record["核对备注"]


def test_alias_assignment_maximizes_complete_one_to_one_mapping():
    rows = [
        {
            "table_color": "深灰",
            "filename": "a.jpg",
            "chart_label": "花灰",
            "confidence": 0.99,
        },
        {
            "table_color": "深灰",
            "filename": "b.jpg",
            "chart_label": "烟灰",
            "confidence": 0.9,
        },
        {
            "table_color": "浅灰",
            "filename": "a.jpg",
            "chart_label": "花灰",
            "confidence": 0.95,
        },
    ]
    result = one_to_one_suggestions(rows, ["深灰", "浅灰"], set())
    assert {(row["table_color"], row["filename"]) for row in result} == {
        ("深灰", "b.jpg"),
        ("浅灰", "a.jpg"),
    }


def test_material_preserves_every_original_character_and_first_component(tmp_path):
    raw = "面料：5％棉、95%聚酯纤维；里料:100%锦纶\n填充物：90%鸭绒 10%羽毛"
    sections = [
        {"target": "fabric", "text": "面料：5％棉、95%聚酯纤维；"},
        {"target": "lining", "text": "里料:100%锦纶\n"},
        {"target": "filling", "text": "填充物：90%鸭绒 10%羽毛"},
    ]
    item = {"id": "0", "first_material": "棉", "sections": sections}
    result = validate_material(raw, item)
    assert result["面料名称(*)"] == "棉"
    assert (
        "".join(
            result[field] for field in ("面料成分(*)", "里料成分(*)", "填充物成分(*)")
        )
        == raw
    )
    with pytest.raises(GeminiError, match="原文"):
        validate_material(
            raw,
            dict(
                item,
                sections=[{"target": "fabric", "text": raw.replace("5％", "5%", 1)}],
            ),
        )
    client = ai()
    with patch.object(
        client, "generate", return_value={"materials": [item]}
    ) as request:
        results = classify_materials(
            [raw, raw],
            client,
            lambda text: "棉",
            [],
            cache=ResultCache("m", tmp_path / "cache"),
        )
        assert (
            request.call_count == 1
            and results[raw]["里料成分(*)"] == sections[1]["text"]
        )


def test_invalid_material_result_falls_back_without_deleting_text(tmp_path):
    raw = "面料:100%锦纶 里料:100%棉"
    client = ai()
    warnings = []
    with patch.object(
        client,
        "generate",
        return_value={
            "materials": [
                {
                    "id": "0",
                    "first_material": "锦纶",
                    "sections": [{"target": "fabric", "text": "100%锦纶"}],
                }
            ]
        },
    ):
        result = classify_materials(
            [raw],
            client,
            lambda text: "锦纶",
            warnings,
            cache=ResultCache("m", tmp_path / "cache"),
        )[raw]
    assert result["面料成分(*)"] + result["里料成分(*)"] == raw
    assert warnings


def master(path):
    book = Workbook()
    s = book.active
    s.title = "新品类层级编码及标费档次"
    for column, value in {
        2: "CAT",
        3: "外套",
        4: "SUB",
        8: "外套",
        10: 199,
        12: "A",
    }.items():
        s.cell(3, column, value)
    s = book.create_sheet("尺码表")
    s.append(["编码", "组"])
    s.append(["SIZE01", "S-3XL"])
    s = book.create_sheet("颜色")
    s.append(["颜色", "编码"])
    s.append(["黑色", "C01"])
    book.save(path)
    return path


def source_book(path, style="新名字", image=None):
    book = Workbook()
    s = book.active
    s.append(["款号", "商品名称", "尺码", "颜色", "上下装", "面料", "成本", "图片"])
    s.append(
        ["1234", style, "S-3XL", "黑色", "上装", "面料:100%锦纶 里料:100%棉", 68, None]
    )
    if image:
        s.add_image(ExcelImage(image), "H2")
    book.save(path)
    return path


def test_optional_floating_picture_extracted_by_column(tmp_path):
    picture = photo(tmp_path / "coat.png", "black")
    source = source_book(tmp_path / "input.xlsx", image=picture)
    book = load_workbook(source)
    images = read_images(source, book.active, 8, [])
    assert 2 in images and Image.open(io.BytesIO(images[2])).size == (64, 64)
    assert not read_images(source, book.active, None, [])
    book.close()


def test_wps_cell_image_formula_extracted(tmp_path):
    source = source_book(tmp_path / "wps.xlsx")
    picture = photo(tmp_path / "photo.png", "blue")
    with zipfile.ZipFile(source) as archive:
        files = {name: archive.read(name) for name in archive.namelist()}
    sheet = files["xl/worksheets/sheet1.xml"].decode()
    sheet = sheet.replace(
        "</row></sheetData>",
        '<c r="H2"><f>DISPIMG("ID_TEST",1)</f></c></row></sheetData>',
    )
    files["xl/worksheets/sheet1.xml"] = sheet.encode()
    files["xl/cellimages.xml"] = (
        b'<root xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><xdr:pic><xdr:nvPicPr><xdr:cNvPr name="ID_TEST"/></xdr:nvPicPr><xdr:blipFill><a:blip r:embed="rId1"/></xdr:blipFill></xdr:pic></root>'
    )
    files["xl/_rels/cellimages.xml.rels"] = (
        b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Target="media/wps.png" Type="image"/></Relationships>'
    )
    files["xl/media/wps.png"] = picture.read_bytes()
    with zipfile.ZipFile(source, "w") as archive:
        for name, data in files.items():
            archive.writestr(name, data)
    book = load_workbook(source, data_only=True)
    assert read_images(source, book.active, 8, [])[2] == picture.read_bytes()
    book.close()


def test_excel_rich_cell_image_metadata_extracted(tmp_path):
    source = source_book(tmp_path / "rich.xlsx")
    picture = photo(tmp_path / "photo.png", "blue")
    with zipfile.ZipFile(source) as archive:
        files = {name: archive.read(name) for name in archive.namelist()}
    sheet = (
        files["xl/worksheets/sheet1.xml"]
        .decode()
        .replace(
            "</row></sheetData>",
            '<c r="H2" vm="1" t="e"><v>#VALUE!</v></c></row></sheetData>',
        )
    )
    files["xl/worksheets/sheet1.xml"] = sheet.encode()
    files["xl/metadata.xml"] = (
        b'<metadata><metadataTypes><metadataType name="XLRICHVALUE"/></metadataTypes><futureMetadata name="XLRICHVALUE"><bk><extLst><rvb i="0"/></extLst></bk></futureMetadata><valueMetadata><bk><rc t="1" v="0"/></bk></valueMetadata></metadata>'
    )
    files["xl/richData/rdrichvaluestructure.xml"] = (
        b'<rvStructures><s><k n="CalcOrigin"/><k n="_rvRel:LocalImageIdentifier"/></s></rvStructures>'
    )
    files["xl/richData/rdrichvalue.xml"] = (
        b'<rvData><rv s="0"><v>99</v><v>0</v></rv></rvData>'
    )
    files["xl/richData/richValueRel.xml"] = (
        b'<richValueRels xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><rel r:id="rId1"/></richValueRels>'
    )
    files["xl/richData/_rels/richValueRel.xml.rels"] = (
        b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Target="../media/rich.png" Type="image"/></Relationships>'
    )
    files["xl/media/rich.png"] = picture.read_bytes()
    with zipfile.ZipFile(source, "w") as archive:
        for name, data in files.items():
            archive.writestr(name, data)
    book = load_workbook(source, data_only=True)
    assert read_images(source, book.active, 8, [])[2] == picture.read_bytes()
    book.close()


@pytest.mark.parametrize("names", [[], ["外套"], ["", "外套"]])
def test_gemini_35_image_schema_has_no_empty_enum_and_unknown_stays_blank(
    tmp_path, names
):
    client = ai()
    client.model = "gemini-3.5-flash-lite"
    picture = photo(tmp_path / "coat.png", "black")
    candidates = dict.fromkeys(names, ("CAT", "SUB"))

    def response(parts, schema):
        for field in ("category_name", "price_name"):
            assert schema["properties"][field]["enum"]
            assert all(value.strip() for value in schema["properties"][field]["enum"])
        return {
            "category_name": schema["properties"]["category_name"]["enum"][-1],
            "price_name": schema["properties"]["price_name"]["enum"][-1],
            "garment_kind": "无法确认",
            "contact": "无法判断",
            "confidence": 0.9,
            "needs_review": False,
            "reason": "没有可匹配的候选",
        }

    with patch.object(client, "generate", side_effect=response):
        result = classify_image(
            client,
            picture.read_bytes(),
            "新名字",
            (candidates, candidates, {}, {}),
            {},
            [],
        )
    assert result["category_name"] == result["price_name"] == ""


def test_image_assisted_generation_maps_only_master_candidates(tmp_path):
    module = load_original("generator")
    picture = photo(tmp_path / "coat.png", "black")
    source = source_book(tmp_path / "input.xlsx", image=picture)
    reference = master(tmp_path / "master.xlsx")
    output = tmp_path / "output.xlsx"
    raw = "面料:100%锦纶 里料:100%棉"
    composition = {
        "materials": [
            {
                "id": "0",
                "first_material": "锦纶",
                "sections": [
                    {"target": "fabric", "text": "面料:100%锦纶 "},
                    {"target": "lining", "text": "里料:100%棉"},
                ],
            }
        ]
    }
    category = {
        "category_name": "外套",
        "price_name": "外套",
        "garment_kind": "梭织外套",
        "contact": "外穿",
        "confidence": 0.95,
        "needs_review": False,
        "reason": "有门襟和外套领型",
    }
    client = ai()
    with patch.object(client, "generate", side_effect=[composition, category]), patch(
        "suite.submission.lookup_standard", return_value={}
    ):
        _, count, _, _ = module.generate(
            source, reference, output, module.App.DEFAULTS, ai_client=client
        )
    book = load_workbook(output)
    values = dict(zip(module.OUTPUT_HEADERS, [cell.value for cell in book.active[2]]))
    assert (
        count == 1 and values["品类编码(*)"] == "CAT" and values["小类编码(*)"] == "SUB"
    )
    assert values["商品名称(*)"].endswith("新名字")
    assert values["面料成分(*)"] + values["里料成分(*)"] == raw
    assert values["执行标准(*)"] is None and values["安全技术级别(*)"] is None
    assert "AI核对" in book.sheetnames
    book.close()


def test_missing_optional_picture_skips_image_ai_and_keeps_empty_fields(tmp_path):
    module = load_original("generator")
    source = source_book(tmp_path / "input.xlsx")
    reference = master(tmp_path / "master.xlsx")
    output = tmp_path / "output.xlsx"
    client = ai()
    with patch("suite.submission.classify_image") as vision, patch.object(
        client, "generate", side_effect=GeminiError("offline")
    ):
        _, count, _, _ = module.generate(
            source, reference, output, module.App.DEFAULTS, ai_client=client
        )
    vision.assert_not_called()
    book = load_workbook(output)
    values = dict(zip(module.OUTPUT_HEADERS, [cell.value for cell in book.active[2]]))
    assert (
        count == 1 and values["品类编码(*)"] is None and values["执行标准(*)"] is None
    )
    assert values["面料成分(*)"] + values["里料成分(*)"] == "面料:100%锦纶 里料:100%棉"
    book.close()


def test_image_fallback_calibrates_existing_rule_standard_and_audits_old_value(
    tmp_path,
):
    module = load_original("generator")
    picture = photo(tmp_path / "coat.png", "black")
    source = source_book(tmp_path / "input.xlsx", style="某款外套", image=picture)
    reference = master(tmp_path / "master.xlsx")
    output = tmp_path / "output.xlsx"
    category = {
        "category_name": "外套",
        "price_name": "外套",
        "garment_kind": "羽绒外套",
        "contact": "外穿",
        "confidence": 0.95,
        "needs_review": False,
        "reason": "模拟可见外套结构",
    }
    verified = {
        "执行标准(*)": "GB/T 14272-2021",
        "安全技术级别(*)": "GB 18401-2010 C类",
        "标准来源": ["https://std.samr.gov.cn/test"],
        "标准依据": "模拟官方来源核对",
    }
    with patch("suite.submission.classify_materials", return_value={}), patch(
        "suite.submission.classify_image", return_value=category
    ), patch("suite.submission.lookup_standard", return_value=verified):
        module.generate(source, reference, output, module.App.DEFAULTS, ai_client=ai())
    book = load_workbook(output)
    values = dict(zip(module.OUTPUT_HEADERS, [cell.value for cell in book.active[2]]))
    audit = dict(
        zip(
            [cell.value for cell in book["AI核对"][1]],
            [cell.value for cell in book["AI核对"][2]],
        )
    )
    assert values["执行标准(*)"] == "GB/T 14272-2021"
    assert audit["原规则执行标准"] == "FZ/T 81007-2022"
    assert audit["输出执行标准"] == "GB/T 14272-2021"
    book.close()


def test_pants_rule_only_knitted_or_woven_pants():
    module = load_original("generator")
    assert module.detect_class("男款针织裤") and module.detect_class("梭织裤")
    assert (
        module.detect_class("牛仔裤") is None and module.detect_class("休闲裤") is None
    )


def test_same_style_unconfirmed_image_is_not_requested_per_row(tmp_path):
    module = load_original("generator")
    picture = photo(tmp_path / "coat.png", "black")
    source = source_book(tmp_path / "input.xlsx", image=picture)
    book = load_workbook(source)
    book.active.append(
        [
            "1234",
            "新名字",
            "S-3XL",
            "黑色",
            "上装",
            "面料:100%锦纶 里料:100%棉",
            68,
            None,
        ]
    )
    book.save(source)
    book.close()
    reference = master(tmp_path / "master.xlsx")
    with patch("suite.submission.classify_materials", return_value={}), patch(
        "suite.submission.classify_image", return_value={}
    ) as classify:
        module.generate(
            source,
            reference,
            tmp_path / "output.xlsx",
            module.App.DEFAULTS,
            ai_client=ai(),
        )
        assert classify.call_count == 1


def test_template_includes_optional_image_column(tmp_path):
    path = tmp_path / "template.xlsx"
    create_input_template(path)
    book = load_workbook(path)
    assert book.active["G1"].value == "图片"
    assert "上下装" not in [cell.value for cell in book.active[1]]
    book.close()


def test_category_only_master_generates_with_size_color_codes_empty(tmp_path):
    module = load_original("generator")
    reference = master(tmp_path / "categories.xlsx")
    book = load_workbook(reference)
    del book["尺码表"]
    del book["颜色"]
    book.save(reference)
    book.close()
    source = source_book(tmp_path / "input.xlsx", style="外套")
    output = tmp_path / "output.xlsx"
    warnings, count, _, _ = module.generate(
        source, reference, output, module.App.DEFAULTS
    )
    book = load_workbook(output)
    values = dict(zip(module.OUTPUT_HEADERS, [cell.value for cell in book.active[2]]))
    assert (
        count == 1 and values["品类编码(*)"] == "CAT" and values["小类编码(*)"] == "SUB"
    )
    assert values["颜色编码(*)"] is None and values["尺码组(*)"] is None
    assert any("未提供【颜色】" in message for message in warnings)
    book.close()


def standard():
    return {
        "execution_standard": "FZ/T 81007-2022",
        "safety_category": "GB 18401-2010 C类",
        "status": "现行",
        "effective_date": "2023-04-01",
        "safety_status": "现行",
        "safety_effective_date": "2012-08-01",
    }


def grounding():
    return {
        "groundingChunks": [
            {
                "web": {
                    "uri": "https://std.samr.gov.cn/fz/product",
                    "title": "FZ/T 81007-2022",
                }
            },
            {
                "web": {
                    "uri": "https://openstd.samr.gov.cn/gb/safety",
                    "title": "GB 18401-2010",
                }
            },
        ],
        "groundingSupports": [{"groundingChunkIndices": [0, 1]}],
    }


def test_standards_require_current_grounded_official_sources():
    result = validate_standard(standard(), grounding())
    assert (
        result["执行标准(*)"] == "FZ/T 81007-2022"
        and result["安全技术级别(*)"] == "GB 18401-2010 C类"
    )
    with pytest.raises(GeminiError):
        validate_standard(standard(), {})
    with pytest.raises(GeminiError):
        validate_standard(dict(standard(), effective_date="2999-01-01"), grounding())
    metadata = grounding()
    metadata["groundingSupports"][0]["groundingChunkIndices"] = [0]
    assert not validate_standard(standard(), metadata)["安全技术级别(*)"]
    metadata = grounding()
    metadata["groundingChunks"][0]["web"]["uri"] = "https://example.com/fiction"
    with pytest.raises(GeminiError):
        validate_standard(standard(), metadata)


def test_side_ai_can_return_two_front_views_without_forcing_a_back(tmp_path):
    items = [
        {"path": str(photo(tmp_path / f"{i}.jpg", "black")), "name": f"{i}.jpg"}
        for i in range(2)
    ]
    reply = {
        "images": [
            {
                "id": str(i),
                "side": "正面",
                "garment_kind": "裤装",
                "confidence": 0.95,
                "needs_review": False,
                "reason": "可见门襟",
            }
            for i in range(2)
        ]
    }
    client = ai()
    module = load_original("joiner")
    with patch.object(client, "generate", return_value=reply):
        recognize_sides(client, items, cache=ResultCache("s", tmp_path / "cache"))
    assign_sides(module, items)
    assert [item["side"] for item in items] == ["正面", "正面"]
    unknown = {"name": "unknown.jpg", "front_score": 0}
    assign_sides(module, [{"name": "front.jpg", "front_score": 100}, unknown])
    assert unknown["side"] == "待确认"
    manual = {
        "name": "front.jpg",
        "manual_side": True,
        "manual_side_value": "背面",
        "side": "待确认",
    }
    assign_sides(module, [manual])
    assert manual["side"] == "背面"
    assert filename_hint("backpack.jpg") is None


def test_pants_local_analysis_uses_fly_and_back_pockets():
    module = load_original("joiner")

    def pants(back):
        image = Image.new("RGB", (200, 320), "white")
        draw = ImageDraw.Draw(image)
        draw.polygon(
            [
                (60, 20),
                (140, 20),
                (150, 300),
                (105, 300),
                (100, 150),
                (95, 300),
                (50, 300),
            ],
            fill="#606070",
        )
        if back:
            draw.rectangle((65, 55, 85, 85), outline="black", width=3)
            draw.rectangle((115, 55, 135, 85), outline="black", width=3)
        else:
            draw.line((100, 25, 100, 80), fill="black", width=3)
        return module.analyze_image(image)

    front, back = pants(False), pants(True)
    assert front["garment_kind"] == back["garment_kind"] == "裤装"
    assert front["front_score"] > back["front_score"]


@pytest.fixture
def tk_root():
    import os, tkinter as tk

    if os.name != "nt" and not os.environ.get("DISPLAY"):
        pytest.skip("Tk needs DISPLAY")
    root = tk.Tk()
    root.withdraw()
    yield root
    root.destroy()
    gc.collect()


def test_matcher_two_color_lists_choose_pid_and_preserve_return_mapping(tk_root):
    module = load_original("matcher")
    app = module.ExcelMatcherApp(tk_root)
    common_ttk_theme(tk_root)
    app.jushuitan_df = pd.DataFrame(
        [
            {"款式编码": style, "颜色": color, "规格": "S", "商品编码": style + color}
            for style, colors in [("A", ["黑色", "白色"]), ("B", ["浅蓝", "深灰色"])]
            for color in colors
        ]
    )
    app.doudian_df = pd.DataFrame(
        [
            {
                "商品id(必填)": pid,
                "商品规格": f"颜色分类:{color};尺码大小:S",
                "商家编码(必填)": "KEEP",
            }
            for pid, colors in [("101", ["黑色", "白色"]), ("202", ["月光蓝", "灰"])]
            for color in colors
        ]
    )
    app.style_list = ["A", "B"]
    app.dd_pid_list = ["101", "202"]
    app.show_style_page(reset=True)
    tk_root.update_idletasks()
    assert not hasattr(app, "import_mapping_table")
    app._choose_reference_pid("101")
    assert app.style_pid_rows["A"][0]["combobox"].get() == "101"
    assert "白色" in app.style_pid_rows["A"][0]["colors_label"].cget("text")
    app._activate_style("B")
    first_card = app._pid_reference_inner.winfo_children()[0]
    assert first_card.winfo_children()[0].winfo_children()[0].cget("text") == "202"
    app._choose_reference_pid("202")
    app.style_to_pid = app._collect_current_style_pid()
    app.show_style_page(reset=False)
    tk_root.update_idletasks()
    assert app._collect_current_style_pid() == {"A": ["101"], "B": ["202"]}
