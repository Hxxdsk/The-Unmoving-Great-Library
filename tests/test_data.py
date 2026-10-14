import json
import threading
import zipfile
from pathlib import Path
from unittest.mock import patch

import pytest
from openpyxl import Workbook
from PIL import Image

from color_tool.data import (ImageMatch, ValidationError, apply_suggestions,
                             export_project, make_plan, read_project, safe_component)
from color_tool.settings import Settings


def make_workbook(root, rows, name="商品资料.xlsx"):
    wb = Workbook()
    ws = wb.active
    ws.append(["供应商款式编码", "颜色", "商品编码"])
    for row in rows:
        ws.append(row)
    path = root / name
    wb.save(path)
    return path


def make_folder(root, code="CYMY9692"):
    folder = root / code
    folder.mkdir()
    Image.new("RGB", (100, 100), "white").save(folder / "颜色.png")
    Image.new("RGB", (100, 100), "black").save(folder / "2.jpg")
    Image.new("RGB", (100, 100), "gray").save(folder / "10.png")
    return folder


@pytest.fixture
def project(tmp_path):
    workbook = make_workbook(tmp_path, [("CYMY9692", "黑色", f"DKU264MWM1246-DY020{i}") for i in range(1, 7)]
                             + [("CYMY9692", "深灰色", "GRAY01")])
    make_folder(tmp_path)
    return read_project(tmp_path, workbook)


def match(filename, color, confidence=.98, review=False):
    return ImageMatch(filename, color, color, confidence, review)


def test_group_and_folder(project):
    assert len(project.styles) == 1
    style = project.styles[0]
    assert style.sku_count == 7
    assert len(style.colors["黑色"]) == 6
    assert style.chart.name == "颜色.png"
    assert [p.name for p in style.images] == ["2.jpg", "10.png"]
    assert not style.issues


def test_bad_erp_dimension(tmp_path):
    path = make_workbook(tmp_path, [("A", "黑色", "S001"), ("A", "黑色", "S002")])
    with zipfile.ZipFile(path) as z:
        entries = {name: z.read(name) for name in z.namelist()}
    import re
    entries["xl/worksheets/sheet1.xml"] = re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1:A1"',
                                                        entries["xl/worksheets/sheet1.xml"])
    with zipfile.ZipFile(path, "w") as z:
        for name, data in entries.items():
            z.writestr(name, data)
    assert read_project(tmp_path, path).sku_count == 2


def test_multistyle_multisheet_and_numeric_zeros(tmp_path):
    wb = Workbook()
    for ws, code, sku in ((wb.active, "A", "a001"), (wb.create_sheet("第二款"), "B", 12)):
        ws.append(["导出说明"])
        ws.append(["颜色", "商品编码", "供应商款式编码"])
        ws.append(["浅灰", sku, code])
        if isinstance(sku, int):
            ws["B3"].number_format = "00000"
    path = tmp_path / "table.xlsx"
    wb.save(path)
    for code in ("A", "B"):
        make_folder(tmp_path, code)
    result = read_project(tmp_path, path)
    assert [s.code for s in result.styles] == ["A", "B"]
    assert result.styles[1].colors["浅灰"] == ["00012"]


def test_duplicate_sku_dedup_and_conflict(tmp_path):
    path = make_workbook(tmp_path, [("A", "黑色", "s1"), ("A", "黑色", "s1")])
    result = read_project(tmp_path, path)
    assert result.sku_count == 1 and len(result.warnings) == 1
    path = make_workbook(tmp_path, [("A", "黑色", "s1"), ("B", "黑色", "s1")])
    with pytest.raises(ValidationError, match="不同款号或颜色"):
        read_project(tmp_path, path)


def test_short_row_reports_missing(tmp_path):
    path = make_workbook(tmp_path, [("A", "黑色")])
    with pytest.raises(ValidationError, match="为空"):
        read_project(tmp_path, path)


def test_missing_and_duplicate_chart(tmp_path):
    path = make_workbook(tmp_path, [("A", "黑色", "s1")])
    assert read_project(tmp_path, path).styles[0].issues
    folder = make_folder(tmp_path, "A")
    Image.new("RGB", (10, 10)).save(folder / "颜色.jpg")
    assert "多张" in read_project(tmp_path, path).styles[0].ai_issues[0]
    assert not read_project(tmp_path, path).styles[0].issues


@pytest.mark.parametrize("name", ["../escape", "a/b", "a\\b", "NUL", "CON.jpg", "x:", "x.", "", "LPT1"])
def test_invalid_filenames(name):
    with pytest.raises(ValidationError):
        safe_component(name)


def test_auto_matching_and_manual_names(project):
    style = project.styles[0]
    style.chart_colors = ["黑色", "花灰色"]
    style.matches = {"2.jpg": match("2.jpg", "黑色"), "10.png": match("10.png", "花灰色")}
    apply_suggestions(style)
    assert style.selections == {"黑色": "2.jpg"}
    style.selections["深灰色"] = "10.png"
    style.manual_colors.add("深灰色")
    apply_suggestions(style)
    assert style.selections["深灰色"] == "10.png"


def test_low_confidence_and_ambiguous_color_not_auto(project):
    style = project.styles[0]
    style.chart_colors = ["黑色"]
    for confidence, review in ((.6, False), (.98, True)):
        style.matches = {"2.jpg": match("2.jpg", "黑色", confidence, review)}
        apply_suggestions(style)
        assert not style.selections
    style.chart_colors = ["黑", "黑色"]
    style.matches = {"2.jpg": match("2.jpg", "黑色")}
    apply_suggestions(style)
    assert not style.selections


def test_color_suffix_equivalence(project):
    style = project.styles[0]
    style.chart_colors = ["深灰"]
    style.matches = {"10.png": match("10.png", "深灰")}
    apply_suggestions(style)
    assert style.selections == {"深灰色": "10.png"}


def test_unmatched_blocks_export(project):
    with pytest.raises(ValidationError, match="还没有选择图片"):
        make_plan(project)


def test_full_copy_same_color_same_bytes_and_repeat(project):
    style = project.styles[0]
    style.selections = {"黑色": "2.jpg", "深灰色": "10.png"}
    plan = make_plan(project)
    original = (style.folder / "2.jpg").read_bytes()
    batch = export_project(project, plan)
    files = list((batch / style.code).glob("*.jpg"))
    assert len(files) == 6
    assert sorted(p.stem for p in files) == style.colors["黑色"]
    assert all(p.read_bytes() == original for p in files)
    assert (batch / style.code / "GRAY01.png").read_bytes() == (style.folder / "10.png").read_bytes()
    assert len(json.loads((batch / "配对记录.json").read_text())['文件']) == 7
    another = export_project(project, plan)
    assert batch != another and batch.exists()
    assert (style.folder / "2.jpg").read_bytes() == original


def test_cancel_and_copy_failure_clean_staging(project):
    style = project.styles[0]
    style.selections = {"黑色": "2.jpg", "深灰色": "10.png"}
    plan = make_plan(project)
    cancel = threading.Event()
    cancel.set()
    with pytest.raises(ValidationError, match="取消"):
        export_project(project, plan, cancel)
    assert not list((project.root / "_商品编码图片").iterdir())
    with patch("color_tool.data.shutil.copy2", side_effect=OSError("disk full")):
        with pytest.raises(OSError, match="disk full"):
            export_project(project, plan)
    assert not list((project.root / "_商品编码图片").iterdir())


def test_settings_model_persists_but_no_key(tmp_path):
    path = tmp_path / "prefs.json"
    settings = Settings(path)
    settings.set("model", "models/gemini-2.5-flash")
    assert Settings(path).get("model") == "models/gemini-2.5-flash"
    with pytest.raises(ValueError):
        settings.set("api_key", "secret")
    assert "secret" not in path.read_text()
