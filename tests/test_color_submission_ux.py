from unittest.mock import patch

import pytest
from PIL import Image
from openpyxl import Workbook, load_workbook
from PySide6.QtWidgets import QApplication
from PySide6.QtTest import QTest

from color_tool.previews import FitImage, ImageChoiceCombo
from suite.legacy import load_original
from suite.submission import create_input_template
from test_color_pairing_recovery import product, reply, item
from color_tool.gemini import GeminiClient
from suite.cache import ResultCache


@pytest.fixture
def qt_app():
    return QApplication.instance() or QApplication([])


def test_full_image_fits_after_resize_and_preserves_aspect(qt_app, tmp_path):
    path = tmp_path / "long.png"
    Image.new("RGB", (1000, 2400), "blue").save(path)
    widget = FitImage(path)
    widget.resize(500, 400)
    widget.show()
    qt_app.processEvents()
    assert widget.pixmap().height() <= 400
    assert widget.pixmap().width() <= 500
    assert abs(widget.pixmap().width() / widget.pixmap().height() - 1000 / 2400) < 0.01
    widget.resize(200, 150)
    qt_app.processEvents()
    assert widget.pixmap().height() <= 150
    widget.close()


def test_hovering_image_choice_shows_preview_without_selecting(qt_app, tmp_path):
    Image.new("RGB", (800, 1000), "purple").save(tmp_path / "紫色.png")
    combo = ImageChoiceCombo(tmp_path)
    combo.addItem("请选择", "")
    combo.addItem("紫色", "紫色.png")
    combo.resize(230, 40)
    combo.show()
    combo.showPopup()
    qt_app.processEvents()
    index = combo.model().index(1, 0)
    QTest.mouseMove(combo.view().viewport(), combo.view().visualRect(index).center())
    qt_app.processEvents()
    assert combo.hover_preview.isVisible()
    assert combo.preview_name.text() == "紫色.png"
    assert combo.preview_image.pixmap().height() > 200
    assert combo.currentIndex() == 0
    combo.hidePopup()
    assert not combo.hover_preview.isVisible()
    combo.close()


def test_overlap_uncertainty_rechecks_only_ambiguous_photo(tmp_path):
    style = product(tmp_path)
    client = GeminiClient("test-key", "models/test")
    initial = reply(
        [
            item("紫图.jpg", "紫色"),
            item("黑图.jpg", "紫色", confidence=0.65, needs_review=True),
        ]
    )
    with patch.object(
        client, "generate", side_effect=[initial, reply([item("黑图.jpg", "黑色")])]
    ) as call:
        client.analyze_style(style, cache=ResultCache("overlap", tmp_path / "cache"))
    assert call.call_count == 2
    prompt = call.call_args_list[1].args[0][0]["text"]
    assert "未遮挡" in prompt and "叠放" in prompt
    assert call.call_args_list[1].args[1]["properties"]["images"]["items"][
        "properties"
    ]["filename"]["enum"] == ["黑图.jpg"]
    assert style.selections == {"紫色": "紫图.jpg", "黑色": "黑图.jpg"}


def generate_sheet(tmp_path, styles, client=None, base_price=199):
    module = load_original("generator")
    master = Workbook()
    s = master.active
    s.title = "新品类层级编码及标费档次"
    for column, value in {
        2: "CAT",
        3: "针织裤",
        4: "SUB",
        8: "针织裤",
        10: base_price,
        12: "A",
    }.items():
        if value is not None:
            s.cell(3, column, value)
    master.save(tmp_path / "master.xlsx")
    create_input_template(tmp_path / "input.xlsx")
    book = load_workbook(tmp_path / "input.xlsx")
    for i, style in enumerate(styles):
        book.active.append([str(i), style, "M", "", "100%棉", 68, None])
    book.save(tmp_path / "input.xlsx")
    book.close()
    with patch("suite.submission.classify_materials", return_value={}):
        module.generate(
            tmp_path / "input.xlsx",
            tmp_path / "master.xlsx",
            tmp_path / "out.xlsx",
            module.App.DEFAULTS,
            ai_client=client,
        )
    book = load_workbook(tmp_path / "out.xlsx")
    rows = [
        dict(zip(module.OUTPUT_HEADERS, [c.value for c in row]))
        for row in book.active.iter_rows(min_row=2)
    ]
    book.close()
    return rows


def test_fleece_matches_base_style_adds_twenty_once_without_affecting_other_prices(
    tmp_path,
):
    rows = generate_sheet(
        tmp_path, ["加绒针织裤", "针织裤（加绒）", "针织裤", "不加绒针织裤"]
    )
    assert [r["标准价(*)"] for r in rows] == [219, 219, 199, 199]
    assert all(r["品类编码(*)"] == "CAT" and r["小类编码(*)"] == "SUB" for r in rows)
    assert all(r["上下装(*)"] == "下装" for r in rows)
    assert rows[0]["商品名称(*)"].endswith("加绒针织裤")
    assert all(r["成本价(*)"] == 78 for r in rows)


def test_missing_base_price_is_not_replaced_with_twenty(tmp_path):
    row = generate_sheet(tmp_path, ["加绒针织裤"], base_price=None)[0]
    assert row["标准价(*)"] is None


def test_position_uses_ai_text_without_image_or_manual_column_and_reuses_style(
    tmp_path, monkeypatch
):
    monkeypatch.setattr("suite.cache.cache_directory", lambda: tmp_path / "cache")

    class Client:
        model = "models/test"
        calls = 0

        def check_cancel(self):
            pass

        def generate(self, parts, schema):
            self.calls += 1
            assert "加绒针织裤" in parts[0]["text"]
            return {"position": "下装", "confidence": 0.97, "reason": "裤装"}

    client = Client()
    rows = generate_sheet(tmp_path, ["加绒针织裤", "加绒针织裤"], client)
    assert client.calls == 1
    assert all(r["上下装(*)"] == "下装" for r in rows)
