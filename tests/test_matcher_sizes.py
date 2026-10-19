import pandas as pd
import pytest

from suite.legacy import load_original
from suite.matcher_adapter import size_key


@pytest.mark.parametrize(
    "raw,expected",
    [
        ("M(建议100-110斤)", "M"),
        ("XL（建议131-150斤）", "XL"),
        ("2XL(加绒)(建议体重)", "2XL"),
        (" M（建议(仅参考)） ", "M"),
        ("xl", "XL"),
        ("L/XL", "L/XL"),
        ("M(未闭合", "M(未闭合"),
        ("（备注）", ""),
    ],
)
def test_size_notes_do_not_change_base_size(raw, expected):
    assert size_key(raw) == expected


def app_for(rows):
    cls = load_original("matcher").ExcelMatcherApp
    app = object.__new__(cls)
    app.jushuitan_df = pd.DataFrame(
        rows, columns=["款式编码", "颜色", "规格", "商品编码"]
    )
    app.style_to_pid = {"A": ["PID"]}
    app.color_mapping = {"A": {"PID": {"黑色": "黑"}}}
    return app


def test_both_sides_normalize_notes_and_sizes_remain_distinct():
    app = app_for(
        [
            ["A", "黑色", "M（基础款）", "M-CODE"],
            ["A", "黑色", "XL", "XL-CODE"],
            ["A", "黑色", "2XL", "2XL-CODE"],
        ]
    )
    lookup = app._build_js_lookup()
    for size, code in [
        ("M(建议100-110斤)", "M-CODE"),
        ("XL（备注）", "XL-CODE"),
        ("2XL", "2XL-CODE"),
    ]:
        assert app._resolve_code("PID", "黑", size, lookup)[:2] == ("ok", code)
    assert app._resolve_code("PID", "黑", "L", lookup)[0] == "missing"
    assert app._resolve_code("PID", "白", "M", lookup)[0] == "unpaired"


def test_conflicting_codes_after_normalization_stay_ambiguous():
    app = app_for(
        [["A", "黑色", "M(薄款)", "SKU-1"], ["A", "黑色", "M（厚款）", "SKU-2"]]
    )
    assert app._resolve_code("PID", "黑", "M", app._build_js_lookup())[0] == "ambiguous"


def test_identical_codes_after_normalization_are_deduplicated():
    app = app_for([["A", "黑色", "M(备注)", "SKU"], ["A", "黑色", "M", "SKU"]])
    assert app._resolve_code("PID", "黑", "M（建议体重）", app._build_js_lookup())[
        :2
    ] == ("ok", "SKU")


def test_export_preserves_spec_notes_and_unpaired_codes(tmp_path):
    from openpyxl import Workbook, load_workbook
    from unittest.mock import patch

    module = load_original("matcher")
    app = object.__new__(module.ExcelMatcherApp)
    app.jushuitan_df = pd.DataFrame(
        [["A", "黑色", "M", "SKU-M"]], columns=["款式编码", "颜色", "规格", "商品编码"]
    )
    app.style_list = ["A"]
    app.style_to_pid = {"A": ["PID"]}
    app.color_mapping = {"A": {"PID": {"黑色": "黑"}}}
    app.block_meta = {("A", "PID"): {"js_colors": ["黑色"], "dd_colors": ["黑", "白"]}}
    app.dd_header_idx = 0
    source, output = tmp_path / "source.xlsx", tmp_path / "output.xlsx"
    book = Workbook()
    book.active.append(["商品id(必填)", "商品规格", "商家编码(必填)"])
    spec = "颜色分类:黑;尺码大小:M（建议100-110斤）"
    book.active.append(["PID", spec, "OLD"])
    book.active.append(["PID", "颜色分类:白;尺码大小:M(备注)", "KEEP"])
    book.save(source)
    app.doudian_path = str(source)
    with patch.object(
        module.filedialog, "asksaveasfilename", return_value=str(output)
    ), patch.object(module.messagebox, "showinfo"), patch.object(
        module.messagebox, "showerror"
    ) as errors:
        app.process_and_save()
    errors.assert_not_called()
    book = load_workbook(output)
    assert book.active["C2"].value == "SKU-M"
    assert book.active["B2"].value == spec
    assert book.active["C3"].value == "KEEP"
    book.close()
    book = load_workbook(source)
    assert book.active["C2"].value == "OLD"
    book.close()
