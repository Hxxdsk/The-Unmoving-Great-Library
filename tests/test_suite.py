import hashlib
import os
import subprocess
import sys
import tkinter as tk
from pathlib import Path
from unittest.mock import MagicMock, patch

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook
from PIL import Image, ImageDraw

from suite.legacy import apply_palette, common_ttk_theme, load_original
from suite.registry import ROOT, TOOLS


@pytest.fixture
def tk_root():
    if os.name != "nt" and not os.environ.get("DISPLAY"):
        pytest.skip("Tk 交互测试需要 DISPLAY；云环境可使用 Xvfb。")
    root = tk.Tk()
    root.withdraw()
    yield root
    root.destroy()


@pytest.mark.parametrize("key", ["matcher", "generator", "joiner"])
def test_original_interfaces_create_with_shared_palette(tk_root, key):
    module = load_original(key)
    apply_palette(module, key)
    classes = {"matcher": "ExcelMatcherApp", "generator": "App", "joiner": "ImageJoinerApp"}
    app = getattr(module, classes[key])(tk_root)
    common_ttk_theme(tk_root)
    tk_root.update_idletasks()
    assert app.root is tk_root
    assert tk_root.winfo_children()


def test_matcher_exports_actual_skus_and_preserves_unpaired(tk_root, tmp_path):
    module = load_original("matcher")
    app = module.ExcelMatcherApp(tk_root)
    app.jushuitan_df = pd.DataFrame([
        {"款式编码": "STYLE-A", "颜色": "黑色", "规格": "S", "商品编码": "SKU-A-S"},
        {"款式编码": "STYLE-A", "颜色": "黑色", "规格": "M", "商品编码": "SKU-A-M"},
        {"款式编码": "STYLE-B", "颜色": "白色", "规格": "S", "商品编码": "SKU-B-S"},
    ])
    path = tmp_path / "抖店.xlsx"
    wb = Workbook()
    wb.active.append(["商品id(必填)", "商品规格", "商家编码(必填)"])
    for color, size in (("黑", "S"), ("黑", "M"), ("白", "S"), ("红", "S")):
        wb.active.append(["PID01", f"颜色分类:{color};尺码大小:{size}", "KEEP"])
    wb.save(path)
    app.doudian_path = str(path)
    app.dd_header_idx = 0
    app.style_list = ["STYLE-A", "STYLE-B"]
    app.style_to_pid = {"STYLE-A": ["PID01"], "STYLE-B": ["PID01"]}
    app.color_mapping = {"STYLE-A": {"PID01": {"黑色": "黑"}}, "STYLE-B": {"PID01": {"白色": "白"}}}
    app.block_meta = {("STYLE-A", "PID01"): {"js_colors": ["黑色"], "dd_colors": ["黑", "白", "红"]},
                      ("STYLE-B", "PID01"): {"js_colors": ["白色"], "dd_colors": ["黑", "白", "红"]}}
    output = tmp_path / "matched.xlsx"
    with patch.object(module.filedialog, "asksaveasfilename", return_value=str(output)), \
            patch.object(module.messagebox, "showinfo"), patch.object(module.messagebox, "showerror") as error:
        app.process_and_save()
    error.assert_not_called()
    loaded = load_workbook(output)
    assert [loaded.active.cell(n, 3).value for n in range(2, 6)] == ["SKU-A-S", "SKU-A-M", "SKU-B-S", "KEEP"]
    loaded.close()
    assert load_workbook(path).active["C2"].value == "KEEP"


def test_generator_writes_submission_with_fabric_and_colors(tmp_path):
    module = load_original("generator")
    master = tmp_path / "主数据.xlsx"
    wb = Workbook()
    category = wb.active
    category.title = "新品类层级编码及标费档次"
    for column, value in {2: "CAT", 3: "外套", 4: "SUB", 8: "外套", 10: 199, 12: "A"}.items():
        category.cell(3, column, value)
    sizes = wb.create_sheet("尺码表")
    sizes.append(["编码", "组"])
    sizes.append(["SIZE01", "S-3XL"])
    colors = wb.create_sheet("颜色")
    colors.append(["颜色", "编码"])
    colors.append(["黑色", "C01"])
    colors.append(["浅灰", "C02"])
    wb.save(master)
    source = tmp_path / "input.xlsx"
    wb = Workbook()
    wb.active.append(module.INPUT_COLUMNS)
    wb.active.append(["1234", "外套", "S-3XL", "黑色;浅灰", "上装", "面料:100%锦纶 里料:100%聚酯纤维", 68])
    wb.save(source)
    target = tmp_path / "submission.xlsx"
    warnings, count, missing, issues = module.generate(str(source), str(master), str(target), module.App.DEFAULTS)
    assert count == 1 and not missing and not issues
    result = load_workbook(target)
    values = dict(zip(module.OUTPUT_HEADERS, [c.value for c in result.active[2]]))
    assert values["开发款号(*)"] == "KQ1234"
    assert values["颜色编码(*)"] == "C01;C02"
    assert values["成本价(*)"] == 78
    assert "100%锦纶" in values["面料成分(*)"]
    fabric_col = module.OUTPUT_HEADERS.index("面料成分(*)") + 1
    assert result.active.cell(2, fabric_col).fill.fgColor.rgb.endswith("FFFF00")
    result.close()


def test_joiner_produces_800_square_front_back_pair(tk_root, tmp_path):
    module = load_original("joiner")
    app = module.ImageJoinerApp(tk_root)
    paths = []
    for name, color in (("正面.png", "#315ab1"), ("背面.png", "#4562a5")):
        image = Image.new("RGB", (200, 400), "white")
        ImageDraw.Draw(image).rectangle((20, 20, 180, 380), fill=color)
        path = tmp_path / name
        image.save(path)
        paths.append(path)
    app.image_list = [{"path": str(path)} for path in paths]
    before = [path.read_bytes() for path in paths]
    with patch.object(module.messagebox, "showinfo"):
        app.process_images(str(tmp_path))
        tk_root.update()
    target = tmp_path / "处理后" / "拼接_1_2.png"
    with Image.open(target) as result:
        assert result.size == (800, 800)
        assert result.getpixel((200, 400)) != result.getpixel((600, 400))
    assert [path.read_bytes() for path in paths] == before


def test_runtime_probe_exercises_current_interpreter():
    args = [sys.executable, str(ROOT / "runtime_probe.py")]
    if os.name != "nt" and not os.environ.get("DISPLAY"):
        args.append("--no-gui")
    result = subprocess.run(args, cwd=ROOT, capture_output=True, timeout=120)
    import json
    report = json.loads(result.stdout)
    assert result.returncode == 0, report
    assert report["ready"] and len(report["checks"]) == 8
    assert report["version"].split(".")[:2] == list(map(str, sys.version_info[:2]))
