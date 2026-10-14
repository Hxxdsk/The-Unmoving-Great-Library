"""兼容完整主数据及只含品类表的主数据；缺失编码不编造。"""

from openpyxl import load_workbook


def load_master(module, path, warnings=None):
    warnings = warnings if warnings is not None else []
    book = load_workbook(path, data_only=True)
    try:
        name = "新品类层级编码及标费档次"
        if name not in book.sheetnames:
            raise KeyError("主数据缺少【新品类层级编码及标费档次】工作表")
        categories, prices, sizes, colors = {}, {}, {}, {}

        def value(sheet, row, column):
            cell = sheet.cell(row, column)
            if cell.value is not None:
                return cell.value
            for area in sheet.merged_cells.ranges:
                if (
                    area.min_row <= row <= area.max_row
                    and area.min_col <= column <= area.max_col
                ):
                    return sheet.cell(area.min_row, area.min_col).value
            return None

        sheet = book[name]
        for row in range(3, sheet.max_row + 1):
            category, code, subcode = (
                value(sheet, row, column) for column in (3, 2, 4)
            )
            if category and code:
                categories.setdefault(
                    module.normalize_key(category),
                    (
                        str(code).strip(),
                        str(subcode).strip() if subcode is not None else "",
                    ),
                )
            price_name, price, fee = (
                value(sheet, row, column) for column in (8, 10, 12)
            )
            if price_name and price is not None:
                prices.setdefault(module.normalize_key(price_name), (price, fee))
        if "尺码表" in book.sheetnames:
            sheet = book["尺码表"]
            for row in range(2, sheet.max_row + 1):
                code, group = value(sheet, row, 1), value(sheet, row, 2)
                if code is not None and group:
                    sizes[module.normalize_key(group)] = str(code).strip()
        else:
            warnings.append("主数据未提供【尺码表】，未匹配尺码编码留空。")
        if "颜色" in book.sheetnames:
            sheet = book["颜色"]
            for row in range(2, sheet.max_row + 1):
                color, code = value(sheet, row, 1), value(sheet, row, 2)
                if color and code:
                    colors[str(color).strip()] = str(code).strip()
        else:
            warnings.append(
                "主数据未提供【颜色】表，未匹配颜色编码留空；补入后可重新生成。"
            )
        return categories, prices, sizes, colors
    finally:
        book.close()
